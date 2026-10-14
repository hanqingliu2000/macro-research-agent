from pathlib import Path

import pytest
from macro_research_agent.workflow.workbook_templates import (
    WORKBOOK_TEMPLATE_VERSION,
    get_workbook_template,
    lint_workbook,
)
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference


def _fixture_workbook(path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    rows = {
        "Analytical Data": [
            ["entity_code", "indicator_code", "period", "value"],
            ["USA", "CPI", "2024", 2.5],
        ],
        "Model Results": [
            ["job_id", "model", "observations", "dropped", "result_hash", "result_json"],
            ["model-1", "ols", 12, 0, "sha256:result", "{}"],
        ],
        "Provenance": [["field", "value"], ["active_dataset_snapshot", "snapshot-1"]],
        "Diagnostics": [["level", "message"], ["diagnostic", "Checks passed."]],
        "Chart Data": [["period", "USA"], ["2024", 2.5]],
    }
    for name, values in rows.items():
        sheet = workbook.create_sheet(name)
        for row in values:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    chart = LineChart()
    chart.add_data(
        Reference(workbook["Chart Data"], min_col=2, min_row=1, max_row=2), titles_from_data=True
    )
    workbook["Chart Data"].add_chart(chart, "D2")
    workbook.properties.description = "publication-v1"
    workbook.save(path)


def test_registered_workbook_template_and_valid_fixture(tmp_path: Path) -> None:
    path = tmp_path / "valid.xlsx"
    _fixture_workbook(path)
    template = get_workbook_template(WORKBOOK_TEMPLATE_VERSION)
    report = lint_workbook(path, require_chart=True)
    assert template.status == "approved"
    assert report.valid is True
    assert report.findings == []
    contract = (
        Path(__file__).resolve().parents[2]
        / "guidelines"
        / "templates"
        / "research-workbook-v1.yaml"
    ).read_text(encoding="utf-8")
    assert f'version: "{WORKBOOK_TEMPLATE_VERSION}"' in contract
    assert 'guideline_version: "publication-v1"' in contract


def test_workbook_lint_returns_rule_addressed_structural_findings(tmp_path: Path) -> None:
    path = tmp_path / "invalid.xlsx"
    _fixture_workbook(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Analytical Data"
    sheet.append(["entity_code", "period", "value"])
    sheet.append(["USA", "2024", "=1+1"])
    workbook.properties.description = "other-v1"
    workbook.save(path)

    report = lint_workbook(path, require_chart=True)
    rule_ids = {item.rule_id for item in report.findings}
    assert report.valid is False
    assert rule_ids >= {
        "WORKBOOK-SHEETS-001",
        "WORKBOOK-COLUMNS-001",
        "WORKBOOK-FREEZE-001",
        "WORKBOOK-FILTER-001",
        "WORKBOOK-FORMULA-001",
        "WORKBOOK-GUIDELINE-001",
        "WORKBOOK-CHART-001",
    }


def test_unknown_workbook_template_fails_closed(tmp_path: Path) -> None:
    path = tmp_path / "valid.xlsx"
    _fixture_workbook(path)
    with pytest.raises(ValueError, match="not registered"):
        lint_workbook(path, template_version="unknown-v1")
