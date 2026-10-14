from __future__ import annotations

import hashlib
import io
import json
import os
import struct
import zipfile
from pathlib import Path
from typing import Any

import httpx
import pandas as pd
import pytest
from macro_research_agent.app import create_app
from macro_research_agent.capabilities.contracts import CapabilityInvocationIntent
from macro_research_agent.config import Settings
from macro_research_agent.workflow.codex_output_reviewer import CodexLocalOutputReviewer
from macro_research_agent.workflow.output_review import (
    DeterministicOutputReviewer,
    OutputReviewFinding,
)
from macro_research_agent.workflow.output_workspace import forecast_delivery_findings
from macro_research_agent.workflow.reporting import (
    DeterministicReportAuthor,
    ReportProviderError,
    ReportRevisionDirectiveV1,
)
from openpyxl import load_workbook
from test_workflow_data_workspace import command_meta
from test_workflow_pipeline import pipeline_client, pipeline_command


async def output_command(
    client: httpx.AsyncClient,
    current: dict[str, Any],
    key: str,
    command: dict[str, Any],
) -> httpx.Response:
    return await client.post(
        "/api/v1/workflow/sessions/planning-1/output/commands",
        json={
            "meta": command_meta(current["projection"]["current_snapshot_hash"], key, "output"),
            "command": command,
        },
    )


class DirectingReportReviewer(DeterministicOutputReviewer):
    provider_profile = "codex-local"
    model_id = "codex-test"
    remote_data_disclosure = False

    def __init__(self) -> None:
        self.calls = 0

    def review(self, request: Any, data_access: Any, render_access: Any) -> Any:
        self.calls += 1
        report = super().review(request, data_access, render_access)
        if self.calls > 1:
            report.semantic_report_status = "pass"
            return report
        artifact = next(
            item
            for item in request.artifacts
            if item.get("format") == "report" and item.get("id") in request.selected_artifact_ids
        )
        artifact_id = str(artifact["id"])
        data_hash = str(request.dataset.grant.content_hash)
        finding = OutputReviewFinding(
            finding_id="finding-report-horizon",
            rule_id="CODEX-REPORT-FORECAST-COVERAGE",
            severity="fail",
            message="The Results section should state the approved horizon.",
            artifact_ids=[artifact_id],
            blocking=True,
            evidence_refs=[data_hash],
        )
        report.findings.append(finding)
        report.status = "fail"
        report.semantic_report_status = "fail"
        report.report_directives = [
            ReportRevisionDirectiveV1(
                directiveId="directive-report-horizon",
                artifactId=artifact_id,
                targetSectionIds=["section-results"],
                issue="State the approved forecast horizon.",
                preserveEvidenceRefs=[data_hash],
                acceptanceCriteria=["State the horizon and its uncertainty."],
            )
        ]
        return report


class FailOnceReportAuthor:
    provider_profile = "deepseek-development"
    model_id = "deepseek-test"
    remote_data_disclosure = True

    def __init__(self) -> None:
        self.calls = 0
        self.delegate = DeterministicReportAuthor()

    def author(self, request: Any, data_access: Any) -> Any:
        self.calls += 1
        if self.calls == 1:
            raise ReportProviderError(
                "DeepSeek returned an invalid report contract after repair.",
                code="invalid_report_contract",
                attempts=2,
                issues=[
                    {
                        "attempt": 2,
                        "path": "$.sections[].heading",
                        "code": "missing_required_sections",
                        "message": "Report draft is missing required sections: Results",
                    }
                ],
            )
        return self.delegate.author(request, data_access)

    def revise(self, request: Any, data_access: Any) -> Any:
        return self.delegate.revise(request, data_access)


def assert_graph_projection_equivalent(response: dict[str, Any]) -> None:
    session = response["session"]
    projection = response["projection"]
    assert session["status"] == projection["status"]
    assert session["activeStage"] == projection["active_stage"]
    assert session["reachedStages"] == projection["reachable_stages"]


def test_forecast_delivery_gate_requires_every_promised_target_and_period() -> None:
    model_jobs = [
        {
            "kind": "var",
            "enabled": True,
            "variables": ["GROWTH", "DEBT"],
            "parameters": {"forecastSteps": 2},
        }
    ]
    incomplete = pd.DataFrame(
        [
            {"variable": "GROWTH", "step": 1},
            {"variable": "GROWTH", "step": 2},
            {"variable": "DEBT", "step": 1},
        ]
    )

    findings = forecast_delivery_findings(model_jobs, incomplete, artifact_id="forecast-1")

    assert len(findings) == 1
    assert findings[0].rule_id == "FORECAST-DELIVERABLE-001"
    assert findings[0].artifact_ids == ["forecast-1"]
    assert "DEBT: periods 2" in findings[0].message
    assert (
        forecast_delivery_findings(
            model_jobs,
            pd.concat([incomplete, pd.DataFrame([{"variable": "DEBT", "step": 2}])]),
        )
        == []
    )


def test_forecast_delivery_ignores_descriptive_only_indicator() -> None:
    jobs = [
        {
            "kind": "var",
            "variables": ["NE.CON.PRVT.ZS", "NE.GDI.FTOT.ZS"],
            "parameters": {"forecastSteps": 2},
        },
        {
            "kind": "descriptive",
            "variables": ["NY.GDP.MKTP.KD.ZG"],
            "parameters": {},
        },
    ]
    complete = pd.DataFrame(
        [
            {"variable": variable, "step": step}
            for variable in ("NE.CON.PRVT.ZS", "NE.GDI.FTOT.ZS")
            for step in (1, 2)
        ]
    )

    assert forecast_delivery_findings(jobs, complete) == []


@pytest.mark.asyncio
async def test_output_artifacts_workbook_manifest_freeze_and_restart(tmp_path: Path) -> None:
    # LRS-OUTPUT-01/10/12: governed artifacts, draft selection and immutable final manifest.
    content = (
        "country,period,indicator,value,note\n"
        "USA,2021,CPI,100,=1+1\nUSA,2022,CPI,108,reviewed\nUSA,2023,CPI,112,reviewed\n"
        "JPN,2021,CPI,100,reviewed\nJPN,2022,CPI,102,reviewed\nJPN,2023,CPI,105,reviewed\n"
    )
    client, current = await pipeline_client(tmp_path, content.encode())
    transport = client._transport  # type: ignore[attr-defined]
    capability_executor = transport.app.state.capability_executor
    capability_executor.execute(
        CapabilityInvocationIntent(
            invocation_id="output-scope-resolution",
            session_id="planning-1",
            stage="plan",
            target_snapshot_hash=current["projection"]["current_snapshot_hash"],
            capability_id="geography.resolve",
            capability_version="1.0.0",
            purpose="Record the controlled geography used by the output package.",
            arguments={"intent": "G7"},
            expected_output_role="geography_resolution",
        )
    )
    added = await pipeline_command(
        client,
        current,
        "output-add-growth",
        {"type": "add_pipeline_operation", "operation": "period_growth"},
    )
    run = await pipeline_command(client, added.json(), "output-run", {"type": "run_pipeline"})
    assert run.status_code == 200, run.text
    body = run.json()
    output = body["session"]["output"]
    assert output["guidelineVersion"] == "publication-v1"
    assert output["workbookTemplateVersion"] == "research-workbook-v1"
    assert output["revisionCapability"] is True
    assert output["artifactSetId"].startswith("artifact-set-")
    assert output["artifactSetHash"].startswith("sha256:")
    artifacts = output["artifacts"]
    assert {item["format"] for item in artifacts} >= {
        "csv",
        "svg",
        "xlsx",
        "python",
        "stata_do",
    }
    assert all(item["status"] == "ready" for item in artifacts)
    data_artifact = next(item for item in artifacts if item["format"] == "csv")
    regenerated_data, regenerated_spec, regenerated_render = (
        transport.app.state.output_workspace_manager.regenerate(
            "planning-1",
            body["session"],
            data_artifact["id"],
            {"source_note": "Source: approved analytical snapshot"},
        )
    )
    assert regenerated_data["id"] != data_artifact["id"]
    assert regenerated_spec["source_note"] == "Source: approved analytical snapshot"
    assert regenerated_render is None

    workbook_artifact = next(item for item in artifacts if item["format"] == "xlsx")
    workbook_response = await client.get(
        f"/api/v1/workflow/sessions/planning-1/output/artifacts/{workbook_artifact['id']}/download"
    )
    assert workbook_response.status_code == 200
    assert (
        "sha256:" + hashlib.sha256(workbook_response.content).hexdigest()
        == workbook_artifact["hash"]
    )
    workbook = load_workbook(io.BytesIO(workbook_response.content), read_only=False)
    assert workbook.sheetnames == [
        "Analytical Data",
        "Model Results",
        "Provenance",
        "Diagnostics",
        "Chart Data",
    ]
    headers = [cell.value for cell in workbook["Analytical Data"][1]]
    assert workbook.properties.description == "publication-v1"
    assert workbook["Analytical Data"].freeze_panes == "A2"
    assert workbook["Analytical Data"].auto_filter.ref
    assert {"entity_code", "indicator_code", "period", "value"}.issubset(headers)
    assert not any(
        cell.data_type == "f" for row in workbook["Analytical Data"].iter_rows() for cell in row
    )
    with zipfile.ZipFile(io.BytesIO(workbook_response.content)) as archive:
        chart_names = [name for name in archive.namelist() if name.startswith("xl/charts/")]
        assert chart_names
        assert b"lineChart" in archive.read(chart_names[0])

    chart_artifact = next(item for item in artifacts if item["format"] == "svg")
    chart_response = await client.get(
        f"/api/v1/workflow/sessions/planning-1/output/artifacts/{chart_artifact['id']}/download"
    )
    assert chart_response.status_code == 200
    assert b"<polyline" in chart_response.content
    assert b"stroke-dasharray" in chart_response.content
    assert b"#d8dee4" in chart_response.content
    assert b"Mock" not in chart_response.content
    chart_spec = output["artifactSpecifications"][chart_artifact["id"]]
    assert chart_spec["guideline_version"] == "publication-v1"
    assert chart_spec["palette"][:3] == ["#0072B2", "#D55E00", "#009E73"]
    assert chart_spec["line_patterns"] == ["solid", "dash", "dot"]

    csv_artifact = next(item for item in artifacts if item["format"] == "csv")
    preview = await client.get(
        f"/api/v1/workflow/sessions/planning-1/output/artifacts/{csv_artifact['id']}"
    )
    assert preview.json()["previewKind"] == "table"
    assert preview.json()["rows"]

    deselected = await output_command(
        client,
        body,
        "deselect-chart",
        {"type": "toggle_artifact_export", "artifactId": chart_artifact["id"]},
    )
    reviewed = await output_command(
        client,
        deselected.json(),
        "review-output",
        {"type": "request_output_review"},
    )
    assert reviewed.status_code == 200, reviewed.text
    assert reviewed.json()["session"]["output"]["reviewReport"]["status"] == "pass"
    frozen = await output_command(
        client, reviewed.json(), "freeze-output", {"type": "finalize_output"}
    )
    replay = await output_command(
        client, reviewed.json(), "freeze-output", {"type": "finalize_output"}
    )
    assert frozen.status_code == 200, frozen.text
    for checkpoint in (
        current,
        added.json(),
        body,
        deselected.json(),
        reviewed.json(),
        frozen.json(),
    ):
        assert_graph_projection_equivalent(checkpoint)
    frozen_output = frozen.json()["session"]["output"]
    assert frozen.json()["session"]["status"] == "completed"
    assert replay.json()["idempotent_replay"] is True
    completed_plan_edit = await client.post(
        "/api/v1/workflow/sessions/planning-1/plan/commands",
        json={
            "meta": command_meta(
                frozen.json()["projection"]["current_snapshot_hash"],
                "edit-completed-plan",
                "output",
            ),
            "command": {
                "type": "save_plan",
                "plan": frozen.json()["session"]["currentPlan"],
            },
        },
    )
    assert completed_plan_edit.status_code == 422
    assert "read-only" in completed_plan_edit.json()["detail"]
    assert (
        next(item for item in frozen_output["artifacts"] if item["id"] == chart_artifact["id"])[
            "status"
        ]
        == "ready"
    )
    manifest = next(item for item in frozen_output["artifacts"] if item["format"] == "json")
    manifest_response = await client.get(
        f"/api/v1/workflow/sessions/planning-1/output/artifacts/{manifest['id']}/download"
    )
    manifest_payload = json.loads(manifest_response.content)
    assert manifest_payload["schemaVersion"] == "mra-reproduction-manifest-v1"
    assert manifest_payload["analyticalBundle"]["id"] == body["session"]["pipeline"]["bundle"]["id"]
    assert manifest_payload["artifactSpecifications"]
    assert manifest_payload["workbookTemplateVersion"] == "research-workbook-v1"
    assert manifest_payload["preflightFindings"] == []
    assert manifest_payload["outputReview"]["status"] == "pass"
    assert (
        manifest_payload["outputReview"]["targetArtifactSetHash"]
        == frozen_output["artifactSetHash"]
    )
    assert chart_artifact["id"] not in {item["id"] for item in manifest_payload["artifacts"]}
    invocation_ids = {item["capability_id"] for item in manifest_payload["capabilityInvocations"]}
    assert invocation_ids == {
        "connector.local_file",
        "geography.resolve",
        "transformation.apply",
        "code.generate",
        "artifact.chart.render",
        "artifact.table.render",
        "output.review.structural",
        "output.review.visual",
    }
    invocation_receipt = next(
        item
        for item in manifest_payload["capabilityInvocations"]
        if item["capability_id"] == "geography.resolve"
    )
    assert invocation_receipt["capability_id"] == "geography.resolve"
    assert invocation_receipt["capability_version"] == "1.0.0"
    assert invocation_receipt["argument_checksum"].startswith("sha256:")
    assert "arguments" not in invocation_receipt
    assert all("arguments" not in item for item in manifest_payload["capabilityInvocations"])
    checkpointed_capabilities = {
        "connector.local_file",
        "transformation.apply",
        "code.generate",
        "artifact.chart.render",
        "artifact.table.render",
    }
    assert all(
        item["checkpoint_ref"]
        for item in manifest_payload["capabilityInvocations"]
        if item["capability_id"] in checkpointed_capabilities
    )
    trace = await client.get("/api/v1/admin/workflow/sessions/planning-1/trace")
    run_event = next(item for item in trace.json() if item["event_type"] == "pipeline.executed")
    freeze_event = next(
        item for item in trace.json() if item["event_type"] == "artifact_set.frozen"
    )
    assert run_event["payload"]["output_artifact_ids"] == [item["id"] for item in artifacts]
    assert freeze_event["payload"]["manifest_artifact_id"] == manifest["id"]
    await client.aclose()

    restarted = create_app(Settings(environment="test", runtime_dir=tmp_path, llm_provider="test"))
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=restarted), base_url="http://test"
    ) as restarted_client:
        recovered = await restarted_client.get(
            f"/api/v1/workflow/sessions/planning-1/output/artifacts/{manifest['id']}/download"
        )
    assert recovered.status_code == 200
    assert recovered.content == manifest_response.content


@pytest.mark.asyncio
async def test_wide_table_revision_selects_columns_and_clears_blocker(tmp_path: Path) -> None:
    columns = [
        "country",
        "period",
        "indicator",
        "value",
        "unit",
        "provider",
        *(f"metadata_{index}" for index in range(20)),
    ]
    content = (
        ",".join(columns)
        + "\n"
        + ",".join(
            ["IND", "2024", "GDP", "100", "Percent", "World Bank", *("x" for _ in range(20))]
        )
    )
    client, current = await pipeline_client(tmp_path, content.encode())
    added = await pipeline_command(
        client,
        current,
        "wide-table-add-log",
        {"type": "add_pipeline_operation", "operation": "log"},
    )
    run = await pipeline_command(client, added.json(), "wide-table-run", {"type": "run_pipeline"})
    assert run.status_code == 200, run.text
    output = run.json()["session"]["output"]
    table = next(item for item in output["artifacts"] if item["format"] == "csv")
    selected = await output_command(
        client,
        run.json(),
        "wide-table-select",
        {"type": "select_artifact", "artifactId": table["id"]},
    )
    reviewed = await output_command(
        client,
        selected.json(),
        "wide-table-review",
        {"type": "request_output_review"},
    )
    report = reviewed.json()["session"]["output"]["reviewReport"]
    wide_findings = [item for item in report["findings"] if item["ruleId"] == "TABLE-COLUMNS-001"]
    identities = {
        (tuple(sorted(item["artifactIds"])), item["ruleId"], item["message"])
        for item in report["findings"]
    }
    assert len(identities) == len(report["findings"])
    assert not wide_findings
    assert len(output["artifactSpecifications"][table["id"]]["columns"]) <= 24

    proposed = await output_command(
        client,
        reviewed.json(),
        "wide-table-proposal",
        {
            "type": "propose_output_revision",
            "request": (
                "Keep a compact publication table with only entity, indicator, period, value, "
                "unit, and provider columns."
            ),
        },
    )
    assert proposed.status_code == 200, proposed.text
    proposal = proposed.json()["session"]["output"]["pendingProposal"]
    assert proposal["patches"][0]["changes"] == {
        "columns": [
            "entity_code",
            "indicator_code",
            "period",
            "value",
            "unit",
            "provider",
        ]
    }
    applied = await output_command(
        client,
        proposed.json(),
        "wide-table-apply",
        {
            "type": "apply_output_revision",
            "acceptedPatchIds": [proposal["patches"][0]["id"]],
        },
    )
    assert applied.status_code == 200, applied.text
    revised = applied.json()["session"]["output"]
    assert "pendingProposal" not in revised
    assert revised["revisionHistory"][-1]["revision"] == 1
    assert not any(
        item["rule_id"] == "TABLE-COLUMNS-001" and table["id"] in item["artifact_ids"]
        for item in revised["preflightFindings"]
    )
    preview = await client.get(
        f"/api/v1/workflow/sessions/planning-1/output/artifacts/{revised['selectedArtifactId']}"
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["columns"] == [
        "entity_code",
        "indicator_code",
        "period",
        "value",
        "unit",
        "provider",
    ]
    await client.aclose()


@pytest.mark.asyncio
async def test_output_splits_multiple_indicators_into_unit_homogeneous_charts(
    tmp_path: Path,
) -> None:
    content = (
        "country,period,indicator,value,unit\n"
        "CHN,2021,GGXWDG,84617439552000,Domestic currency\n"
        "CHN,2022,GGXWDG,95334932033000,Domestic currency\n"
        "CHN,2023,GGXWDG,108000000000000,Domestic currency\n"
        "CHN,2021,GGXWDG_NGDP,71.5,Percent of GDP\n"
        "CHN,2022,GGXWDG_NGDP,77.1,Percent of GDP\n"
        "CHN,2023,GGXWDG_NGDP,83.6,Percent of GDP\n"
    )
    client, current = await pipeline_client(tmp_path, content.encode())
    model_added = await pipeline_command(
        client,
        current,
        "multi-chart-model",
        {"type": "add_model_job", "model": "descriptive"},
    )
    model = model_added.json()["session"]["pipeline"]["models"][0]
    configured = await pipeline_command(
        client,
        model_added.json(),
        "multi-chart-configure",
        {
            "type": "update_model_specification",
            "modelId": model["id"],
            "variables": ["GGXWDG", "GGXWDG_NGDP"],
        },
    )
    run = await pipeline_command(
        client, configured.json(), "multi-chart-run", {"type": "run_pipeline"}
    )
    assert run.status_code == 200, run.text
    output = run.json()["session"]["output"]
    charts = [item for item in output["artifacts"] if item["format"] == "svg"]
    assert {item["name"] for item in charts} == {
        "analytical_chart_ggxwdg.svg",
        "analytical_chart_ggxwdg_ngdp.svg",
    }
    chart_specs = [output["artifactSpecifications"][item["id"]] for item in charts]
    assert {item["indicator_code"] for item in chart_specs} == {"GGXWDG", "GGXWDG_NGDP"}
    assert {item["unit"] for item in chart_specs} == {"Domestic currency", "Percent of GDP"}
    for chart in charts:
        chart_response = await client.get(
            f"/api/v1/workflow/sessions/planning-1/output/artifacts/{chart['id']}/download"
        )
        specification = output["artifactSpecifications"][chart["id"]]
        assert specification["indicator_code"].encode() in chart_response.content
        assert f"Unit: {specification['unit']}".encode() in chart_response.content
        assert chart_response.content.count(b'font-size="20"') >= 1
    debt_chart = next(item for item in charts if item["name"].endswith("_ggxwdg.svg"))
    debt_response = await client.get(
        f"/api/v1/workflow/sessions/planning-1/output/artifacts/{debt_chart['id']}/download"
    )
    assert b"T</text>" in debt_response.content

    workbook_artifact = next(item for item in output["artifacts"] if item["format"] == "xlsx")
    workbook_response = await client.get(
        f"/api/v1/workflow/sessions/planning-1/output/artifacts/{workbook_artifact['id']}/download"
    )
    workbook = load_workbook(io.BytesIO(workbook_response.content), read_only=False)
    assert {"Chart Data", "Chart Data 2"}.issubset(workbook.sheetnames)

    reviewed = await output_command(
        client, run.json(), "multi-chart-review", {"type": "request_output_review"}
    )
    revision = await pipeline_command(
        client,
        reviewed.json(),
        "multi-chart-start-analysis-revision",
        {"type": "start_pipeline_revision"},
    )
    edited_after_review = await pipeline_command(
        client,
        revision.json(),
        "multi-chart-edit-after-review",
        {
            "type": "update_model_specification",
            "modelId": model["id"],
            "variables": ["GGXWDG_NGDP", "GGXWDG"],
        },
    )
    rerun = await pipeline_command(
        client,
        edited_after_review.json(),
        "multi-chart-rerun-after-review",
        {"type": "run_pipeline"},
    )
    await client.aclose()
    assert reviewed.status_code == 200, reviewed.text
    assert reviewed.json()["session"]["output"]["reviewReport"]["status"] == "pass"
    assert edited_after_review.status_code == 200, edited_after_review.text
    assert edited_after_review.json()["session"]["activeStage"] == "transformation"
    assert rerun.status_code == 200, rerun.text
    assert rerun.json()["session"]["activeStage"] == "output"


@pytest.mark.asyncio
async def test_output_review_typed_revision_and_selective_regeneration(tmp_path: Path) -> None:
    # LRS-OUTPUT-03/04/06/08/09: complete-data/render review and approved typed revision.
    content = (
        "country,period,indicator,value\n"
        "USA,2021,CPI,100\nUSA,2022,CPI,108\nUSA,2023,CPI,112\n"
        "JPN,2021,CPI,100\nJPN,2022,CPI,102\nJPN,2023,CPI,105\n"
    )
    client, current = await pipeline_client(tmp_path, content.encode())
    added = await pipeline_command(
        client,
        current,
        "review-add-growth",
        {"type": "add_pipeline_operation", "operation": "period_growth"},
    )
    run = await pipeline_command(client, added.json(), "review-run", {"type": "run_pipeline"})
    assert run.status_code == 200, run.text
    output = run.json()["session"]["output"]
    chart = next(item for item in output["artifacts"] if item["format"] == "svg")
    review_ref = output["reviewRenderRefs"][chart["id"]]
    review_render = await client.get(
        f"/api/v1/workflow/sessions/planning-1/output/artifacts/{review_ref['id']}/download"
    )
    assert review_render.status_code == 200
    assert review_render.content.startswith(b"\x89PNG\r\n\x1a\n")
    width, height = struct.unpack(">II", review_render.content[16:24])
    assert (width, height) == (900, 500)
    assert "sha256:" + hashlib.sha256(review_render.content).hexdigest() == review_ref["hash"]

    blocked = await output_command(
        client, run.json(), "finish-before-review", {"type": "finalize_output"}
    )
    assert blocked.status_code == 422
    assert "Run Output review" in blocked.json()["detail"]

    reviewed = await output_command(
        client, run.json(), "run-output-review", {"type": "request_output_review"}
    )
    assert reviewed.status_code == 200, reviewed.text
    report = reviewed.json()["session"]["output"]["reviewReport"]
    assert report["status"] == "pass"
    assert report["providerProfile"] == "deterministic-local"
    assert report["inspectedRowCount"] == 6
    assert report["reviewRenderHashes"][chart["id"]] == review_ref["hash"]

    read_only_artifact = next(item for item in output["artifacts"] if item["format"] == "python")
    selected_report = await output_command(
        client,
        reviewed.json(),
        "select-read-only-artifact-for-revision",
        {"type": "select_artifact", "artifactId": read_only_artifact["id"]},
    )
    blocked_report_revision = await output_command(
        client,
        selected_report.json(),
        "reject-read-only-revision-scope",
        {"type": "propose_output_revision", "request": "Rewrite this code"},
    )
    assert blocked_report_revision.status_code == 422
    assert "limited to chart and table" in blocked_report_revision.json()["detail"]

    selected = await output_command(
        client,
        selected_report.json(),
        "select-chart-for-revision",
        {"type": "select_artifact", "artifactId": chart["id"]},
    )
    proposed = await output_command(
        client,
        selected.json(),
        "propose-chart-revision",
        {"type": "propose_output_revision", "request": "Short reviewed chart title"},
    )
    assert proposed.status_code == 200, proposed.text
    proposal = proposed.json()["session"]["output"]["pendingProposal"]
    assert proposal["targetArtifactSetHash"] == output["artifactSetHash"]
    assert proposal["patches"][0]["changes"] == {"title": "Short reviewed chart title"}
    unaffected_before = {
        item["id"]: item["hash"] for item in output["artifacts"] if item["id"] != chart["id"]
    }
    applied = await output_command(
        client,
        proposed.json(),
        "apply-chart-revision",
        {
            "type": "apply_output_revision",
            "acceptedPatchIds": [proposal["patches"][0]["id"]],
        },
    )
    assert applied.status_code == 200, applied.text
    revised_output = applied.json()["session"]["output"]
    revised_chart = next(item for item in revised_output["artifacts"] if item["format"] == "svg")
    assert revised_chart["id"] != chart["id"]
    assert revised_chart["hash"] != chart["hash"]
    assert {
        item["id"]: item["hash"]
        for item in revised_output["artifacts"]
        if item["id"] != revised_chart["id"]
    } == unaffected_before
    assert (
        revised_output["reviewReport"]["targetArtifactSetHash"] == revised_output["artifactSetHash"]
    )
    assert revised_output["reviewIteration"] == 2
    assert revised_output["selectedArtifactId"] == revised_chart["id"]
    revision = revised_output["revisionHistory"][0]
    assert revision["revision"] == 1
    assert revision["request"] == "Short reviewed chart title"
    assert revision["reviewIteration"] == 2
    assert revision["reviewReportId"] == revised_output["reviewReport"]["id"]
    assert revision["artifacts"] == [
        {
            "previousArtifactId": chart["id"],
            "artifactId": revised_chart["id"],
            "name": revised_chart["name"],
            "format": "svg",
            "previousHash": chart["hash"],
            "hash": revised_chart["hash"],
        }
    ]
    comparison = await client.get(
        f"/api/v1/workflow/sessions/planning-1/output/artifacts/{revised_chart['id']}"
        f"/compare/{chart['id']}"
    )
    assert comparison.status_code == 200, comparison.text
    comparison_body = comparison.json()
    assert comparison_body["status"] == "presentation_changed"
    assert comparison_body["hash_changed"] is True
    assert comparison_body["value_comparison"]["status"] == "unchanged"
    assert "title" in {item["field"] for item in comparison_body["specification_changes"]}
    old_chart = await client.get(
        f"/api/v1/workflow/sessions/planning-1/output/artifacts/{chart['id']}/download"
    )
    assert old_chart.status_code == 200

    second_proposal = await output_command(
        client,
        applied.json(),
        "propose-final-chart-revision",
        {"type": "propose_output_revision", "request": "Final reviewed chart title"},
    )
    assert second_proposal.status_code == 200, second_proposal.text
    second_patch = second_proposal.json()["session"]["output"]["pendingProposal"]["patches"][0]
    second_applied = await output_command(
        client,
        second_proposal.json(),
        "apply-final-chart-revision",
        {"type": "apply_output_revision", "acceptedPatchIds": [second_patch["id"]]},
    )
    assert second_applied.status_code == 200, second_applied.text
    bounded_output = second_applied.json()["session"]["output"]
    assert bounded_output["reviewIteration"] == 3
    assert bounded_output["maxReviewIterations"] == 3
    assert len(bounded_output["revisionHistory"]) == 2
    blocked_fourth_review = await output_command(
        client,
        second_applied.json(),
        "block-fourth-output-review",
        {"type": "request_output_review"},
    )
    assert blocked_fourth_review.status_code == 422
    assert "iteration limit" in blocked_fourth_review.json()["detail"]
    blocked_third_revision = await output_command(
        client,
        second_applied.json(),
        "block-third-output-revision",
        {"type": "propose_output_revision", "request": "One revision too many"},
    )
    assert blocked_third_revision.status_code == 422
    assert "iteration limit" in blocked_third_revision.json()["detail"]

    graph = await client.get("/api/v1/admin/workflow/sessions/planning-1/graph")
    assert any(
        item["authority"] == "review" and item["kind"] == "visual_artifact_review"
        for item in graph.json()["nodes"]
    )
    assert any(
        item["authority"] == "proposal" and item["kind"] == "output_revision"
        for item in graph.json()["nodes"]
    )
    await client.aclose()


@pytest.mark.skipif(
    os.getenv("MRA_LIVE_CODEX_OUTPUT_REVIEW") != "1",
    reason="Requires explicit opt-in and a locally authenticated Codex account.",
)
@pytest.mark.asyncio
async def test_live_codex_local_reviews_a_generated_output_artifact_set(tmp_path: Path) -> None:
    content = (
        "country,period,indicator,value,unit\n"
        "CHN,2022,CPI,100,Index\nCHN,2023,CPI,100.2,Index\nCHN,2024,CPI,100.4,Index\n"
        "USA,2022,CPI,100,Index\nUSA,2023,CPI,104.1,Index\nUSA,2024,CPI,107.1,Index\n"
    )
    client, current = await pipeline_client(tmp_path, content.encode())
    added = await pipeline_command(
        client,
        current,
        "live-codex-add-growth",
        {"type": "add_pipeline_operation", "operation": "period_growth"},
    )
    run = await pipeline_command(client, added.json(), "live-codex-run", {"type": "run_pipeline"})
    assert run.status_code == 200, run.text
    transport = client._transport  # type: ignore[attr-defined]
    reviewer = CodexLocalOutputReviewer(
        working_root=tmp_path / ".codex-output-review",
        timeout_seconds=120,
    )
    transport.app.state.output_reviewer = reviewer
    transport.app.state.workflow_output_service.output_reviewer = reviewer

    reviewed = await output_command(
        client,
        run.json(),
        "live-codex-output-review",
        {"type": "request_output_review"},
    )

    assert reviewed.status_code == 200, reviewed.text
    report = reviewed.json()["session"]["output"]["reviewReport"]
    assert report["providerProfile"] == "codex-local"
    assert "CODEX-LOCAL-REVIEW-UNAVAILABLE" not in {
        finding["ruleId"] for finding in report["findings"]
    }
    assert "VISUAL-RUBRIC-PLACEHOLDER" in {finding["ruleId"] for finding in report["findings"]}
    assert report["visualRubricVersion"] == "academic-publication-placeholder-v1"
    assert report["visualRubricHash"].startswith("sha256:")
    assert report["reviewedArtifactIds"]
    print(
        json.dumps(
            {
                "status": report["status"],
                "reviewed_artifact_count": len(report["reviewedArtifactIds"]),
                "findings": [
                    {
                        "rule_id": finding["ruleId"],
                        "severity": finding["severity"],
                        "message": finding["message"],
                    }
                    for finding in report["findings"]
                ],
            },
            ensure_ascii=False,
        )
    )
    await client.aclose()


@pytest.mark.asyncio
async def test_report_author_failure_preserves_analysis_and_retries_output_only(
    tmp_path: Path,
) -> None:
    content = (
        "country,period,indicator,value\n"
        "USA,2021,CPI,100\nUSA,2022,CPI,108\nUSA,2023,CPI,112\n"
        "JPN,2021,CPI,100\nJPN,2022,CPI,102\nJPN,2023,CPI,105\n"
    )
    client, current = await pipeline_client(
        tmp_path,
        content.encode(),
        request="Produce a research report comparing inflation outcomes.",
    )
    transport = client._transport  # type: ignore[attr-defined]
    manager = transport.app.state.output_workspace_manager
    author = FailOnceReportAuthor()
    manager.report_author = author
    added = await pipeline_command(
        client,
        current,
        "recoverable-report-add-growth",
        {"type": "add_pipeline_operation", "operation": "period_growth"},
    )
    model_added = await pipeline_command(
        client,
        added.json(),
        "recoverable-report-add-model",
        {"type": "add_model_job", "model": "descriptive"},
    )
    model = model_added.json()["session"]["pipeline"]["models"][0]
    configured = await pipeline_command(
        client,
        model_added.json(),
        "recoverable-report-configure-model",
        {
            "type": "update_model_specification",
            "modelId": model["id"],
            "variables": ["CPI"],
        },
    )

    failed = await pipeline_command(
        client,
        configured.json(),
        "recoverable-report-run",
        {"type": "run_pipeline"},
    )

    assert failed.status_code == 200, failed.text
    failed_body = failed.json()
    failed_session = failed_body["session"]
    failed_pipeline = failed_session["pipeline"]
    assert failed_session["status"] == "failed"
    assert failed_session["activeStage"] == "output"
    assert failed_pipeline["runStatus"] == "completed"
    assert failed_pipeline["bundle"]["id"]
    assert len(failed_pipeline["modelResults"]) == 1
    assert failed_pipeline["outputGeneration"] == {
        "status": "failed",
        "attempt": 1,
        "bundleId": failed_pipeline["bundle"]["id"],
        "retryable": True,
        "lastAttemptAt": failed_session["updatedAt"],
        "error": {
            "code": "invalid_report_contract",
            "message": "DeepSeek returned an invalid report contract after repair.",
            "issues": [
                {
                    "attempt": 2,
                    "path": "$.sections[].heading",
                    "code": "missing_required_sections",
                    "message": "Report draft is missing required sections: Results",
                }
            ],
        },
    }
    assert failed_session["output"]["artifacts"] == []
    assert not list(
        (manager.workflow_root / "sessions/planning-1/output").glob("artifact-*")
    )
    failed_node = next(
        item for item in failed_session["nodes"] if item["kind"] == "output_generation"
    )
    assert failed_node["status"] == "failed"
    assert failed_node["errorCode"] == "invalid_report_contract"
    bundle_id = failed_pipeline["bundle"]["id"]
    model_result_ids = [item["id"] for item in failed_pipeline.get("modelResults") or []]

    retried = await pipeline_command(
        client,
        failed_body,
        "recoverable-report-retry",
        {"type": "retry_output_generation"},
    )
    replay = await pipeline_command(
        client,
        failed_body,
        "recoverable-report-retry",
        {"type": "retry_output_generation"},
    )

    assert retried.status_code == 200, retried.text
    retried_body = retried.json()
    retried_session = retried_body["session"]
    assert retried_session["status"] == "output_review"
    assert retried_session["pipeline"]["bundle"]["id"] == bundle_id
    assert [
        item["id"] for item in retried_session["pipeline"].get("modelResults") or []
    ] == model_result_ids
    assert retried_session["pipeline"]["outputGeneration"]["status"] == "completed"
    assert retried_session["pipeline"]["outputGeneration"]["attempt"] == 2
    assert retried_session["output"]["reportDrafts"]
    assert author.calls == 2
    assert replay.status_code == 200
    assert replay.json()["idempotent_replay"] is True
    assert replay.json()["session"]["output"]["artifactSetId"] == retried_session["output"][
        "artifactSetId"
    ]
    trace = await client.get("/api/v1/admin/workflow/sessions/planning-1/trace")
    failures = [
        item for item in trace.json() if item["event_type"] == "output.generation_failed"
    ]
    assert len(failures) == 1
    assert failures[0]["payload"]["issues"][0]["path"] == "$.sections[].heading"
    await client.aclose()


@pytest.mark.asyncio
async def test_publication_report_is_spec_bound_and_content_reviewed(tmp_path: Path) -> None:
    content = (
        "country,period,indicator,value\n"
        "USA,2021,CPI,100\nUSA,2022,CPI,108\nUSA,2023,CPI,112\n"
        "JPN,2021,CPI,100\nJPN,2022,CPI,102\nJPN,2023,CPI,105\n"
    )
    client, current = await pipeline_client(
        tmp_path,
        content.encode(),
        request="Produce a research report comparing inflation outcomes.",
    )
    added = await pipeline_command(
        client,
        current,
        "report-add-growth",
        {"type": "add_pipeline_operation", "operation": "period_growth"},
    )
    run = await pipeline_command(client, added.json(), "report-run", {"type": "run_pipeline"})
    assert run.status_code == 200, run.text
    output = run.json()["session"]["output"]
    report_artifact = next(item for item in output["artifacts"] if item["format"] == "report")
    report_spec = output["artifactSpecifications"][report_artifact["id"]]
    assert report_spec["schema_version"] == "mra-report-spec-v2"
    assert report_spec["observation_count"] == 6
    assert report_spec["draft_hash"].startswith("sha256:")
    assert report_spec["evidence_ledger_hash"].startswith("sha256:")
    assert output["reportDrafts"][report_artifact["id"]]["claims"]
    assert output["evidenceLedgers"][report_artifact["id"]][0]["evidenceRefs"]
    assert output["reportAuthorRequests"][report_artifact["id"]]["dataset"]["initial_page"] is None
    assert report_spec["required_sections"] == [
        "Scope",
        "Data and method",
        "Results",
        "Diagnostics and limitations",
        "Reproducibility",
    ]
    reviewed = await output_command(
        client, run.json(), "review-report", {"type": "request_output_review"}
    )
    assert reviewed.status_code == 200, reviewed.text
    review = reviewed.json()["session"]["output"]["reviewReport"]
    assert review["findings"] == []
    assert review["status"] == "pass", review
    assert review["semanticReportStatus"] == "pass"
    assert not any(item["ruleId"].startswith("REPORT-") for item in review["findings"])
    report_download = await client.get(
        f"/api/v1/workflow/sessions/planning-1/output/artifacts/{report_artifact['id']}/download"
    )
    assert report_download.status_code == 200
    assert b"## Reproducibility" in report_download.content
    graph = await client.get("/api/v1/admin/workflow/sessions/planning-1/graph")
    assert any(
        item["kind"] == "semantic_report_review" and item["authority"] == "review"
        for item in graph.json()["nodes"]
    )
    await client.aclose()


@pytest.mark.asyncio
async def test_approved_report_directive_rewrites_only_report_then_reviews_again(
    tmp_path: Path,
) -> None:
    content = (
        "country,period,indicator,value\n"
        "USA,2021,CPI,100\nUSA,2022,CPI,108\nUSA,2023,CPI,112\n"
        "JPN,2021,CPI,100\nJPN,2022,CPI,102\nJPN,2023,CPI,105\n"
    )
    client, current = await pipeline_client(
        tmp_path,
        content.encode(),
        request="Produce a research report comparing inflation outcomes.",
    )
    added = await pipeline_command(
        client,
        current,
        "report-directive-add-growth",
        {"type": "add_pipeline_operation", "operation": "period_growth"},
    )
    run = await pipeline_command(
        client, added.json(), "report-directive-run", {"type": "run_pipeline"}
    )
    assert run.status_code == 200, run.text
    transport = client._transport  # type: ignore[attr-defined]
    reviewer = DirectingReportReviewer()
    transport.app.state.output_reviewer = reviewer
    transport.app.state.workflow_output_service.output_reviewer = reviewer

    reviewed = await output_command(
        client,
        run.json(),
        "report-directive-review",
        {"type": "request_output_review"},
    )
    assert reviewed.status_code == 200, reviewed.text
    reviewed_output = reviewed.json()["session"]["output"]
    proposal = reviewed_output["pendingProposal"]
    assert proposal["patches"] == []
    assert proposal["reportDirectives"][0]["targetSectionIds"] == ["section-results"]
    old_artifacts = {item["id"]: item for item in reviewed_output["artifacts"]}
    old_report = next(item for item in old_artifacts.values() if item["format"] == "report")

    revised = await output_command(
        client,
        reviewed.json(),
        "apply-report-directive",
        {
            "type": "apply_output_revision",
            "acceptedPatchIds": [],
            "acceptedDirectiveIds": ["directive-report-horizon"],
        },
    )

    assert revised.status_code == 200, revised.text
    output = revised.json()["session"]["output"]
    new_report = next(item for item in output["artifacts"] if item["format"] == "report")
    assert new_report["id"] != old_report["id"]
    assert new_report["hash"] != old_report["hash"]
    assert {
        item["id"]: item["hash"] for item in output["artifacts"] if item["format"] != "report"
    } == {item["id"]: item["hash"] for item in old_artifacts.values() if item["format"] != "report"}
    assert output["reviewIteration"] == 2
    assert output["reportSemanticStatus"] == "pass"
    assert output["reviewReport"]["semanticReportStatus"] == "pass"
    assert output["revisionHistory"][0]["acceptedDirectiveIds"] == ["directive-report-horizon"]
    assert "pendingProposal" not in output
    reopened = await pipeline_command(
        client,
        revised.json(),
        "reopen-after-report-review",
        {"type": "start_pipeline_revision"},
    )
    assert reopened.status_code == 200, reopened.text
    graph = await client.get("/api/v1/admin/workflow/sessions/planning-1/graph")
    semantic_nodes = [
        item for item in graph.json()["nodes"] if item["kind"] == "semantic_report_review"
    ]
    assert semantic_nodes and all(item["lifecycle"] == "stale" for item in semantic_nodes)
    await client.aclose()
