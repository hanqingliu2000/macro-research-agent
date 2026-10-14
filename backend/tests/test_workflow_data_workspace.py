from __future__ import annotations

import io
import json
from pathlib import Path

import httpx
import pytest
from macro_research_agent.app import create_app
from macro_research_agent.config import Settings
from openpyxl import Workbook, load_workbook
from test_workflow_planning import meta, planning_session


def command_meta(snapshot_hash: str, key: str, stage: str = "data") -> dict[str, str]:
    value = meta(snapshot_hash, key)
    value["expected_stage"] = stage
    return value


async def approved_data_client(
    tmp_path: Path,
    content: bytes,
    filename: str = "macro.csv",
    request: str | None = None,
    include_plan_analysis: bool = True,
) -> tuple[httpx.AsyncClient, dict[str, object]]:
    app = create_app(Settings(environment="test", runtime_dir=tmp_path, llm_provider="test"))
    client = httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test")
    initial_session = planning_session()
    if request is not None:
        initial_session["request"] = request
    created = await client.post("/api/v1/workflow/sessions", json={"session": initial_session})
    current_hash = created.json()["projection"]["current_snapshot_hash"]
    await client.post(
        "/api/v1/workflow/sessions/planning-1/sources/uploads",
        files={"file": (filename, content)},
    )
    drafted = await client.post(
        "/api/v1/workflow/sessions/planning-1/plan/commands",
        json={
            "meta": meta(current_hash, "data-draft"),
            "command": {
                "type": "submit_clarification",
                "answers": {
                    "geography_scope": "G7",
                    "period_frequency": "2019-2024 quarterly",
                },
            },
        },
    )
    if not include_plan_analysis:
        drafted = await client.post(
            "/api/v1/workflow/sessions/planning-1/plan/commands",
            json={
                "meta": meta(
                    drafted.json()["projection"]["current_snapshot_hash"],
                    "clear-plan-analysis",
                ),
                "command": {
                    "type": "save_plan",
                    "changes": {"transformationIntents": [], "modelIntents": []},
                },
            },
        )
        assert drafted.status_code == 200, drafted.text
    confirmed = await client.post(
        "/api/v1/workflow/sessions/planning-1/plan/commands",
        json={
            "meta": meta(drafted.json()["projection"]["current_snapshot_hash"], "data-confirm"),
            "command": {"type": "confirm_geography", "resolutionId": "geo-primary"},
        },
    )
    approved = await client.post(
        "/api/v1/workflow/sessions/planning-1/plan/commands",
        json={
            "meta": meta(confirmed.json()["projection"]["current_snapshot_hash"], "data-approve"),
            "command": {"type": "plan_next"},
        },
    )
    assert approved.status_code == 200, approved.text
    return client, approved.json()


@pytest.mark.asyncio
async def test_local_file_normalizes_to_canonical_snapshot_and_paginates(tmp_path: Path) -> None:
    client, approved = await approved_data_client(
        tmp_path,
        b"country,period,indicator,value,note\n"
        b"United States,2024-Q1,CPI,3.1,a\n"
        b"JPN,2024-Q1,CPI,2.2,b\n",
    )
    workspace = await client.get("/api/v1/workflow/sessions/planning-1/data")
    preview = await client.get("/api/v1/workflow/sessions/planning-1/data/preview?offset=1&limit=1")
    await client.aclose()

    data = workspace.json()
    assert approved["session"]["activeStage"] == "data"
    assert [column["id"] for column in data["columns"][:6]] == [
        "series_id",
        "entity_code",
        "indicator_code",
        "frequency",
        "period",
        "value",
    ]
    assert data["quality"]["observations"] == 2
    assert data["sourceRuns"][0]["connectorId"] == "local_file"
    assert data["entityMappings"]["United States"] == "USA"
    assert data["provenance"][0]["rawChecksum"].startswith("sha256:")
    assert preview.json()["total_rows"] == 2
    assert preview.json()["rows"][0]["entity_code"] == "JPN"
    assert Path(
        next(
            (tmp_path / "workflow" / "sessions" / "planning-1" / "data" / "snapshots").glob(
                "*.parquet"
            )
        )
    ).exists()
    restarted = create_app(Settings(environment="test", runtime_dir=tmp_path, llm_provider="test"))
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=restarted), base_url="http://test"
    ) as restarted_client:
        recovered = await restarted_client.get(
            "/api/v1/workflow/sessions/planning-1/data/preview?offset=0&limit=2"
        )
    assert recovered.status_code == 200
    assert [row["entity_code"] for row in recovered.json()["rows"]] == ["USA", "JPN"]


@pytest.mark.asyncio
async def test_wide_macro_file_is_melted_to_long_form(tmp_path: Path) -> None:
    client, approved = await approved_data_client(
        tmp_path,
        b"country,year,gdp,inflation\nUSA,2024,29000,3.1\n",
    )
    await client.aclose()
    data = approved["session"]["data"]
    assert data["quality"]["observations"] == 2
    indicators = {row["indicator_code"] for row in data["snapshots"][0]["rows"]}
    assert indicators == {"GDP", "INFLATION"}


@pytest.mark.asyncio
async def test_missing_share_counts_observation_values_not_optional_metadata(
    tmp_path: Path,
) -> None:
    client, approved = await approved_data_client(
        tmp_path,
        b"country,period,indicator,value,note\n"
        b"USA,2024,CPI,3.1,\n"
        b"JPN,2024,CPI,,\n",
    )
    await client.aclose()

    quality = approved["session"]["data"]["quality"]
    assert quality["observations"] == 2
    assert quality["missingShare"] == 50.0
    assert quality["warning"].startswith("1 observation value(s) are missing.")


@pytest.mark.asyncio
async def test_json_and_xlsx_use_the_same_canonical_loader_contract(tmp_path: Path) -> None:
    json_client, json_result = await approved_data_client(
        tmp_path / "json-runtime",
        b'{"observations":[{"country":"USA","period":"2024","indicator":"GDP","value":29000}]}',
        "macro.json",
    )
    await json_client.aclose()

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["country", "period", "indicator", "value"])
    sheet.append(["Japan", "2024", "GDP", 4200])
    payload = io.BytesIO()
    workbook.save(payload)
    xlsx_client, xlsx_result = await approved_data_client(
        tmp_path / "xlsx-runtime", payload.getvalue(), "macro.xlsx"
    )
    await xlsx_client.aclose()

    assert json_result["session"]["data"]["quality"]["observations"] == 1
    assert json_result["session"]["data"]["snapshots"][0]["rows"][0]["entity_code"] == "USA"
    assert xlsx_result["session"]["data"]["quality"]["observations"] == 1
    assert xlsx_result["session"]["data"]["snapshots"][0]["rows"][0]["entity_code"] == "JPN"


@pytest.mark.asyncio
async def test_unrecognized_source_schema_is_mapped_before_loading(tmp_path: Path) -> None:
    client, approved = await approved_data_client(
        tmp_path,
        b"economy_name,time_key,metric_name,amount_raw\nUSA,2024-Q1,CPI,3.1\n",
    )
    draft = approved["session"]["data"]["sourceMappingDraft"]
    assert approved["session"]["data"]["activeSnapshotId"] == ""
    current_hash = approved["projection"]["current_snapshot_hash"]
    mappings = {
        "economy_name": "entity_code",
        "time_key": "period",
        "metric_name": "indicator_code",
        "amount_raw": "value",
    }
    latest = approved
    for index, (column, role) in enumerate(mappings.items()):
        response = await client.post(
            "/api/v1/workflow/sessions/planning-1/data/commands",
            json={
                "meta": command_meta(current_hash, f"source-map-{index}"),
                "command": {
                    "type": "map_source_column",
                    "sourceColumn": column,
                    "role": role,
                },
            },
        )
        assert response.status_code == 200
        latest = response.json()
        current_hash = latest["projection"]["current_snapshot_hash"]
    loaded = await client.post(
        "/api/v1/workflow/sessions/planning-1/data/commands",
        json={
            "meta": command_meta(current_hash, "source-validate"),
            "command": {"type": "validate_source_mapping"},
        },
    )
    await client.aclose()

    assert draft["state"] == "mapping"
    assert loaded.status_code == 200
    assert loaded.json()["session"]["data"]["activeSnapshotId"]
    assert loaded.json()["session"]["data"]["quality"]["observations"] == 1


@pytest.mark.asyncio
async def test_typed_deduplicate_undo_and_data_approval(tmp_path: Path) -> None:
    client, approved = await approved_data_client(
        tmp_path,
        b"country,period,indicator,value\nUSA,2024-Q1,CPI,3.0\nUSA,2024-Q1,CPI,3.1\nJPN,2024-Q1,CPI,2.2\n",
    )
    initial_hash = approved["projection"]["current_snapshot_hash"]
    blocked = await client.post(
        "/api/v1/workflow/sessions/planning-1/data/commands",
        json={
            "meta": command_meta(initial_hash, "blocked-data-next"),
            "command": {"type": "data_next"},
        },
    )
    deduplicated = await client.post(
        "/api/v1/workflow/sessions/planning-1/data/commands",
        json={
            "meta": command_meta(initial_hash, "deduplicate"),
            "command": {"type": "apply_data_operation", "operation": "deduplicate"},
        },
    )
    replay = await client.post(
        "/api/v1/workflow/sessions/planning-1/data/commands",
        json={
            "meta": command_meta(initial_hash, "deduplicate"),
            "command": {"type": "apply_data_operation", "operation": "deduplicate"},
        },
    )
    stale = await client.post(
        "/api/v1/workflow/sessions/planning-1/data/commands",
        json={
            "meta": command_meta(initial_hash, "stale-data-operation"),
            "command": {"type": "apply_data_operation", "operation": "sort"},
        },
    )
    undone = await client.post(
        "/api/v1/workflow/sessions/planning-1/data/commands",
        json={
            "meta": command_meta(
                deduplicated.json()["projection"]["current_snapshot_hash"], "undo"
            ),
            "command": {"type": "undo_data_operation"},
        },
    )
    redone = await client.post(
        "/api/v1/workflow/sessions/planning-1/data/commands",
        json={
            "meta": command_meta(
                undone.json()["projection"]["current_snapshot_hash"], "deduplicate-again"
            ),
            "command": {"type": "apply_data_operation", "operation": "deduplicate"},
        },
    )
    next_stage = await client.post(
        "/api/v1/workflow/sessions/planning-1/data/commands",
        json={
            "meta": command_meta(redone.json()["projection"]["current_snapshot_hash"], "data-next"),
            "command": {"type": "data_next"},
        },
    )
    reapproved = await client.post(
        "/api/v1/workflow/sessions/planning-1/data/commands",
        json={
            "meta": command_meta(
                next_stage.json()["projection"]["current_snapshot_hash"],
                "data-next-again",
                stage="transformation",
            ),
            "command": {"type": "data_next"},
        },
    )
    await client.aclose()

    assert deduplicated.status_code == 200
    assert blocked.status_code == 422
    assert "duplicate" in blocked.json()["detail"].lower()
    assert replay.json()["idempotent_replay"] is True
    assert len(replay.json()["session"]["data"]["snapshots"]) == 2
    assert stale.status_code == 409
    assert deduplicated.json()["session"]["data"]["quality"]["duplicateKeys"] == 0
    assert deduplicated.json()["session"]["data"]["quality"]["observations"] == 2
    assert undone.json()["session"]["data"]["quality"]["duplicateKeys"] == 2
    assert len(undone.json()["session"]["data"]["operations"]) == 0
    assert next_stage.status_code == 200
    assert next_stage.json()["session"]["activeStage"] == "transformation"
    assert next_stage.json()["session"]["data"]["approvedSnapshotId"]
    assert next_stage.json()["session"]["data"]["approvedDatasetRef"]["schemaHash"].startswith(
        "sha256:"
    )
    pipeline = next_stage.json()["session"]["pipeline"]
    assert pipeline["draftSource"] == "approved_plan"
    assert pipeline["sourcePlanSnapshotId"] == next_stage.json()["session"][
        "approvedPlanSnapshotId"
    ]
    assert [item["kind"] for item in pipeline["operations"]] == ["yoy_growth"]
    assert pipeline["operations"][0]["target"]["indicators"] == ["CPI"]
    assert [item["kind"] for item in pipeline["models"]] == ["descriptive"]
    assert pipeline["models"][0]["variables"] == ["CPI"]
    assert pipeline["models"][0]["status"] == "not_run"
    assert reapproved.status_code == 200, reapproved.text
    assert [item["id"] for item in reapproved.json()["session"]["pipeline"]["operations"]] == [
        item["id"] for item in pipeline["operations"]
    ]


@pytest.mark.asyncio
async def test_all_registered_hygiene_operations_are_typed_and_snapshot_backed(
    tmp_path: Path,
) -> None:
    client, current = await approved_data_client(
        tmp_path,
        b"country,period,indicator,value,note\n"
        b" usa ,2024-Q2,cpi,3.1,a\n"
        b"USA,2024-Q1,CPI,,b\n"
        b"USA,2024-Q1,CPI,,b\n",
    )
    operations = [
        "missing_value_rule",
        "filter_missing",
        "map_codes",
        "cast_types",
        "sort",
        "rename_columns",
        "drop_columns",
        "keep_columns",
        "deduplicate",
    ]
    for index, operation in enumerate(operations):
        response = await client.post(
            "/api/v1/workflow/sessions/planning-1/data/commands",
            json={
                "meta": command_meta(
                    current["projection"]["current_snapshot_hash"], f"operation-{index}"
                ),
                "command": {"type": "apply_data_operation", "operation": operation},
            },
        )
        assert response.status_code == 200, response.text
        current = response.json()

    rejected_operation = await client.post(
        "/api/v1/workflow/sessions/planning-1/data/commands",
        json={
            "meta": command_meta(
                current["projection"]["current_snapshot_hash"], "unknown-operation"
            ),
            "command": {"type": "apply_data_operation", "operation": "run_python"},
        },
    )
    rejected_parameters = await client.post(
        "/api/v1/workflow/sessions/planning-1/data/commands",
        json={
            "meta": command_meta(
                current["projection"]["current_snapshot_hash"], "invalid-parameters"
            ),
            "command": {
                "type": "apply_data_operation",
                "operation": "sort",
                "parameters": {"columns": ["not_a_column"]},
            },
        },
    )
    await client.aclose()

    data = current["session"]["data"]
    assert len(data["operations"]) == len(operations)
    assert len(data["snapshots"]) == len(operations) + 1
    assert data["quality"]["observations"] == 1
    assert [column["id"] for column in data["columns"]] == [
        "series_id",
        "entity_code",
        "indicator_code",
        "frequency",
        "period",
        "value",
    ]
    assert rejected_operation.status_code == 422
    assert rejected_operation.json()["detail"] == "Unknown Data operation."
    assert rejected_parameters.status_code == 422
    assert rejected_parameters.json()["detail"] == (
        "Invalid parameters for the selected Data operation."
    )


@pytest.mark.asyncio
async def test_roundtrip_export_import_and_commit_preserve_contract(tmp_path: Path) -> None:
    client, approved = await approved_data_client(
        tmp_path,
        b"country,period,indicator,value\nUSA,2024-Q1,CPI,3.1\nJPN,2024-Q1,CPI,2.2\n",
    )
    exported = await client.post(
        "/api/v1/workflow/sessions/planning-1/data/commands",
        json={
            "meta": command_meta(
                approved["projection"]["current_snapshot_hash"], "roundtrip-export"
            ),
            "command": {"type": "prepare_round_trip_export"},
        },
    )
    package = exported.json()["session"]["data"]["lastExportedPackage"]
    download = await client.get(package["downloadUrl"])
    workbook = load_workbook(io.BytesIO(download.content))
    assert workbook.sheetnames == ["Data", "Schema", "Provenance", "Instructions", "_MRA_MANIFEST"]
    assert workbook["_MRA_MANIFEST"].sheet_state == "hidden"
    assert workbook["_MRA_MANIFEST"].protection.sheet is True
    assert workbook.security.lockStructure is True
    workbook["Data"]["F1"] = "cleaned_value"
    workbook["Data"]["F2"] = 3.5
    edited = io.BytesIO()
    workbook.save(edited)
    import_meta = command_meta(
        exported.json()["projection"]["current_snapshot_hash"], "roundtrip-import"
    )
    imported = await client.post(
        "/api/v1/workflow/sessions/planning-1/data/roundtrip/imports",
        data={"meta_json": json.dumps(import_meta)},
        files={"file": ("cleaned.xlsx", edited.getvalue())},
    )
    mapped = await client.post(
        "/api/v1/workflow/sessions/planning-1/data/commands",
        json={
            "meta": command_meta(
                imported.json()["projection"]["current_snapshot_hash"], "roundtrip-map"
            ),
            "command": {
                "type": "map_round_trip_column",
                "sourceColumn": "cleaned_value",
                "role": "value",
            },
        },
    )
    validated = await client.post(
        "/api/v1/workflow/sessions/planning-1/data/commands",
        json={
            "meta": command_meta(
                mapped.json()["projection"]["current_snapshot_hash"], "roundtrip-validate"
            ),
            "command": {"type": "validate_round_trip"},
        },
    )
    committed = await client.post(
        "/api/v1/workflow/sessions/planning-1/data/commands",
        json={
            "meta": command_meta(
                validated.json()["projection"]["current_snapshot_hash"], "roundtrip-commit"
            ),
            "command": {"type": "commit_round_trip"},
        },
    )
    await client.aclose()

    assert download.status_code == 200
    assert imported.status_code == 200, imported.text
    assert imported.json()["session"]["data"]["roundTripDraft"]["state"] == "mapping"
    assert validated.json()["session"]["data"]["roundTripDraft"]["state"] == "validated"
    assert committed.status_code == 200
    assert committed.json()["session"]["data"]["operations"] == []
    assert (
        committed.json()["session"]["data"]["activeRoundTripManifest"]["manifestVerified"] is True
    )
    active_id = committed.json()["session"]["data"]["activeSnapshotId"]
    active = next(
        item for item in committed.json()["session"]["data"]["snapshots"] if item["id"] == active_id
    )
    assert active["sourceKind"] == "external_round_trip"
    assert active["rows"][0]["value"] == 3.5


@pytest.mark.asyncio
async def test_roundtrip_rejects_formulas_and_manifest_tampering(tmp_path: Path) -> None:
    client, approved = await approved_data_client(
        tmp_path,
        b"country,period,indicator,value\nUSA,2024-Q1,CPI,3.1\n",
    )
    exported = await client.post(
        "/api/v1/workflow/sessions/planning-1/data/commands",
        json={
            "meta": command_meta(approved["projection"]["current_snapshot_hash"], "invalid-export"),
            "command": {"type": "prepare_round_trip_export"},
        },
    )
    package = exported.json()["session"]["data"]["lastExportedPackage"]
    download = await client.get(package["downloadUrl"])
    workbook = load_workbook(io.BytesIO(download.content))
    workbook["Data"]["F2"] = "=1+1"
    workbook["_MRA_MANIFEST"]["B2"] = "tampered"
    edited = io.BytesIO()
    workbook.save(edited)
    imported = await client.post(
        "/api/v1/workflow/sessions/planning-1/data/roundtrip/imports",
        data={
            "meta_json": json.dumps(
                command_meta(
                    exported.json()["projection"]["current_snapshot_hash"],
                    "invalid-import",
                )
            )
        },
        files={"file": ("invalid.xlsx", edited.getvalue())},
    )
    await client.aclose()
    draft = imported.json()["session"]["data"]["roundTripDraft"]
    assert imported.status_code == 200
    assert draft["state"] == "invalid"
    assert draft["manifestVerified"] is False
    assert draft["formulasMaterialized"] is False
