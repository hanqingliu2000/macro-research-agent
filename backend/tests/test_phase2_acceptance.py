from __future__ import annotations

import io
import json
from pathlib import Path
from typing import Any

import httpx
import pytest
from macro_research_agent.app import create_app
from macro_research_agent.config import Settings
from openpyxl import load_workbook

REQUEST = (
    "Compare post-pandemic inflation persistence and monetary-policy responses in the "
    "United States, Japan and the euro area from 2018Q1 through 2025Q4. Produce charts, "
    "tables, clean data, reproducible code and a research report."
)


def _fixture_csv() -> bytes:
    rows = ["country,period,indicator,value,analyst_note"]
    for entity_index, entity in enumerate(("USA", "JPN", "DEU")):
        for year in range(2018, 2026):
            for quarter in range(1, 5):
                index = (year - 2018) * 4 + quarter - 1
                cpi = 100 + entity_index * 2 + index * (0.7 + entity_index * 0.1)
                rate = 1.0 + entity_index * 0.1 + index * 0.08
                rows.append(f"{entity},{year}-Q{quarter},CPI,{cpi:.3f},reviewed")
                rows.append(f"{entity},{year}-Q{quarter},POLICY_RATE,{rate:.3f},reviewed")
    rows.append("USA,2025-Q4,CPI,,prompt-like text is inert")
    return ("\n".join(rows) + "\n").encode()


def _meta(current: dict[str, Any], key: str) -> dict[str, str]:
    return {
        "command_id": f"phase2-{key}",
        "idempotency_key": f"phase2-{key}",
        "session_id": current["session"]["id"],
        "expected_stage": current["projection"]["active_stage"],
        "expected_snapshot_hash": current["projection"]["current_snapshot_hash"],
        "actor": "user",
        "requested_at": "2026-09-06T12:00:00Z",
    }


async def _command(
    client: httpx.AsyncClient,
    current: dict[str, Any],
    stage: str,
    key: str,
    command: dict[str, Any],
) -> httpx.Response:
    session_id = current["session"]["id"]
    return await client.post(
        f"/api/v1/workflow/sessions/{session_id}/{stage}/commands",
        json={"meta": _meta(current, key), "command": command},
    )


@pytest.mark.asyncio
async def test_phase2_longitudinal_local_acceptance(tmp_path: Path) -> None:
    """Near-real T0/T1 journey derived from LRS-001, without deferred live providers."""

    settings = Settings(environment="test", runtime_dir=tmp_path, llm_provider="test")
    app = create_app(settings)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        created = await client.post("/api/v1/workflow/sessions", json={"request": REQUEST})
        assert created.status_code == 201, created.text
        current = created.json()
        session_id = current["session"]["id"]
        assert current["session"]["clarification"]["round"] == 1

        upload = await client.post(
            f"/api/v1/workflow/sessions/{session_id}/sources/uploads",
            files={"file": ("lrs001_macro.csv", _fixture_csv(), "text/csv")},
        )
        assert upload.status_code == 201, upload.text
        drafted = await _command(
            client,
            current,
            "plan",
            "clarify",
            {
                "type": "submit_clarification",
                "answers": {
                    "geography_scope": "United States, Japan and euro area comparison",
                    "period_frequency": "2018-2025 quarterly",
                },
            },
        )
        assert drafted.status_code == 200, drafted.text
        current = drafted.json()
        proposed = await _command(
            client,
            current,
            "plan",
            "plan-proposal",
            {"type": "propose_plan_patch", "request": "State the descriptive limits clearly."},
        )
        assert proposed.status_code == 200, proposed.text
        current = proposed.json()
        accepted = await _command(
            client,
            current,
            "plan",
            "plan-accept",
            {"type": "apply_plan_patch", "acceptedFields": ["assumptions"]},
        )
        assert accepted.status_code == 200, accepted.text
        current = accepted.json()
        geography = await _command(
            client,
            current,
            "plan",
            "geography",
            {"type": "confirm_geography", "resolutionId": "geo-primary"},
        )
        assert geography.status_code == 200, geography.text
        current = geography.json()
        approved = await _command(client, current, "plan", "plan-next", {"type": "plan_next"})
        assert approved.status_code == 200, approved.text
        current = approved.json()
        assert current["projection"]["active_stage"] == "data"
        assert current["session"]["data"]["quality"]["missingShare"] > 0

        data_proposal = await _command(
            client,
            current,
            "data",
            "data-proposal",
            {"type": "propose_data_assistance", "request": "Clean missing rows and sort."},
        )
        assert data_proposal.status_code == 200, data_proposal.text
        current = data_proposal.json()
        data_items = current["session"]["data"]["pendingAssistantProposal"]["operations"]
        accepted_item_ids = [item["id"] for item in data_items]
        cleaned = await _command(
            client,
            current,
            "data",
            "data-accept",
            {"type": "apply_data_assistance", "acceptedItemIds": accepted_item_ids},
        )
        assert cleaned.status_code == 200, cleaned.text
        current = cleaned.json()
        assert current["session"]["data"]["quality"]["missingShare"] == 0

        exported = await _command(
            client,
            current,
            "data",
            "roundtrip-export",
            {"type": "prepare_round_trip_export"},
        )
        assert exported.status_code == 200, exported.text
        current = exported.json()
        package = current["session"]["data"]["lastExportedPackage"]
        workbook_response = await client.get(package["downloadUrl"])
        assert workbook_response.status_code == 200

        invalid_book = load_workbook(io.BytesIO(workbook_response.content))
        value_column = next(
            cell.column for cell in invalid_book["Data"][1] if cell.value == "value"
        )
        invalid_book["Data"].cell(row=2, column=value_column, value="=1+1")
        invalid_bytes = io.BytesIO()
        invalid_book.save(invalid_bytes)
        invalid_import = await client.post(
            f"/api/v1/workflow/sessions/{session_id}/data/roundtrip/imports",
            data={"meta_json": json.dumps(_meta(current, "roundtrip-invalid"))},
            files={"file": ("invalid.xlsx", invalid_bytes.getvalue())},
        )
        assert invalid_import.status_code == 200, invalid_import.text
        assert invalid_import.json()["session"]["data"]["roundTripDraft"]["state"] == "invalid"
        current = invalid_import.json()

        corrected_book = load_workbook(io.BytesIO(workbook_response.content))
        period_column = next(
            cell.column for cell in corrected_book["Data"][1] if cell.value == "period"
        )
        corrected_book["Data"].cell(row=1, column=period_column, value="quarter")
        corrected_bytes = io.BytesIO()
        corrected_book.save(corrected_bytes)
        corrected_import = await client.post(
            f"/api/v1/workflow/sessions/{session_id}/data/roundtrip/imports",
            data={"meta_json": json.dumps(_meta(current, "roundtrip-corrected"))},
            files={"file": ("corrected.xlsx", corrected_bytes.getvalue())},
        )
        assert corrected_import.status_code == 200, corrected_import.text
        current = corrected_import.json()
        assert current["session"]["data"]["roundTripDraft"]["state"] == "mapping"
        mapped = await _command(
            client,
            current,
            "data",
            "roundtrip-map",
            {"type": "map_round_trip_column", "sourceColumn": "quarter", "role": "period"},
        )
        assert mapped.status_code == 200, mapped.text
        current = mapped.json()
        validated = await _command(
            client, current, "data", "roundtrip-validate", {"type": "validate_round_trip"}
        )
        assert validated.status_code == 200, validated.text
        current = validated.json()
        committed = await _command(
            client, current, "data", "roundtrip-commit", {"type": "commit_round_trip"}
        )
        assert committed.status_code == 200, committed.text
        current = committed.json()
        assert current["session"]["data"]["activeRoundTripManifest"]["manifestVerified"]
        data_next = await _command(client, current, "data", "data-next", {"type": "data_next"})
        assert data_next.status_code == 200, data_next.text
        current = data_next.json()

        analysis_proposal = await _command(
            client,
            current,
            "pipeline",
            "analysis-proposal",
            {"type": "propose_analysis_assistance", "request": "Compare inflation persistence."},
        )
        assert analysis_proposal.status_code == 200, analysis_proposal.text
        current = analysis_proposal.json()
        proposal = current["session"]["pipeline"]["pendingAssistantProposal"]
        proposal_ids = [item["id"] for item in [*proposal["operations"], *proposal["models"]]]
        analysis_applied = await _command(
            client,
            current,
            "pipeline",
            "analysis-accept",
            {"type": "apply_analysis_assistance", "acceptedItemIds": proposal_ids},
        )
        assert analysis_applied.status_code == 200, analysis_applied.text
        current = analysis_applied.json()
        run = await _command(client, current, "pipeline", "pipeline-run", {"type": "run_pipeline"})
        assert run.status_code == 200, run.text
        current = run.json()
        assert current["projection"]["active_stage"] == "output"
        output = current["session"]["output"]
        assert {item["format"] for item in output["artifacts"]} >= {
            "csv",
            "xlsx",
            "svg",
            "python",
            "report",
        }
        blocked_finish = await _command(
            client, current, "output", "premature-finish", {"type": "finalize_output"}
        )
        assert blocked_finish.status_code == 422

        reviewed = await _command(
            client, current, "output", "output-review", {"type": "request_output_review"}
        )
        assert reviewed.status_code == 200, reviewed.text
        current = reviewed.json()
        assert current["session"]["output"]["reviewReport"]["status"] in {"pass", "warning"}
        chart = next(
            item for item in current["session"]["output"]["artifacts"] if item["format"] == "svg"
        )
        selected = await _command(
            client,
            current,
            "output",
            "select-chart",
            {"type": "select_artifact", "artifactId": chart["id"]},
        )
        current = selected.json()
        revision = await _command(
            client,
            current,
            "output",
            "revision-proposal",
            {"type": "propose_output_revision", "request": "Inflation persistence comparison"},
        )
        assert revision.status_code == 200, revision.text
        current = revision.json()
        patch = current["session"]["output"]["pendingProposal"]["patches"][0]
        revised = await _command(
            client,
            current,
            "output",
            "revision-apply",
            {"type": "apply_output_revision", "acceptedPatchIds": [patch["id"]]},
        )
        assert revised.status_code == 200, revised.text
        current = revised.json()
        revised_chart = next(
            item for item in current["session"]["output"]["artifacts"] if item["format"] == "svg"
        )
        comparison = await client.get(
            f"/api/v1/workflow/sessions/{session_id}/output/artifacts/"
            f"{revised_chart['id']}/compare/{chart['id']}"
        )
        assert comparison.status_code == 200, comparison.text
        assert comparison.json()["status"] == "presentation_changed"
        assert comparison.json()["value_comparison"]["status"] == "unchanged"

        review_status = current["session"]["output"]["reviewReport"]["status"]
        finalized = await _command(
            client,
            current,
            "output",
            "finish",
            {"type": "finalize_output", "acceptWarnings": review_status == "warning"},
        )
        assert finalized.status_code == 200, finalized.text
        finished = finalized.json()
        assert finished["session"]["status"] == "completed"
        frozen_hashes = {
            item["id"]: item["hash"] for item in finished["session"]["output"]["artifacts"]
        }
        archive = await client.get(f"/api/v1/workflow/sessions/{session_id}/archive")
        assert archive.status_code == 200

    restarted = create_app(settings)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=restarted), base_url="http://test"
    ) as client:
        resumed = await client.get(f"/api/v1/workflow/sessions/{session_id}")
        assert resumed.status_code == 200
        assert {
            item["id"]: item["hash"] for item in resumed.json()["session"]["output"]["artifacts"]
        } == frozen_hashes
        trace = await client.get(f"/api/v1/admin/workflow/sessions/{session_id}/trace")
        sequences = [item["sequence"] for item in trace.json()]
        assert sequences == list(range(1, len(sequences) + 1))
        graph = await client.get(f"/api/v1/admin/workflow/sessions/{session_id}/graph")
        assert graph.status_code == 200
        assert {item["gate"] for item in graph.json()["checkpoints"]} == {
            "research_plan",
            "pipeline_execution",
            "final_outputs",
        }

        imported = await client.post(
            "/api/v1/workflow/sessions/import",
            files={"file": ("lrs001.mra.zip", archive.content, "application/zip")},
        )
        assert imported.status_code == 201, imported.text
        assert imported.json()["session"]["lineage"]["relationship"] == "import"
        forked = await client.post(
            f"/api/v1/workflow/sessions/{session_id}/fork",
            json={"title": "Inflation persistence — continued"},
        )
        assert forked.status_code == 201, forked.text
        assert forked.json()["session"]["status"] == "output_review"
        assert forked.json()["session"]["lineage"]["sourceSessionId"] == session_id
        source_after_fork = await client.get(f"/api/v1/workflow/sessions/{session_id}")
        assert source_after_fork.json()["session"]["status"] == "completed"
        assert {
            item["id"]: item["hash"]
            for item in source_after_fork.json()["session"]["output"]["artifacts"]
        } == frozen_hashes
