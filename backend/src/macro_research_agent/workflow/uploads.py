from __future__ import annotations

import csv
import hashlib
import json
import shutil
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import BinaryIO, Literal
from uuid import uuid4
from zipfile import BadZipFile

from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.utils.exceptions import InvalidFileException  # type: ignore[import-untyped]
from pydantic import BaseModel, Field


class UploadValidationError(ValueError):
    pass


class UploadedColumnSummary(BaseModel):
    name: str
    inferred_type: Literal["string", "number", "boolean", "date", "mixed", "empty"]
    non_null_count: int = Field(ge=0)
    null_count: int = Field(ge=0)


class UploadedTableSummary(BaseModel):
    name: str
    row_count: int = Field(ge=0)
    column_count: int = Field(ge=0)
    columns: list[UploadedColumnSummary]


class UploadedFileRecord(BaseModel):
    upload_id: str
    session_id: str
    filename: str
    media_type: str
    byte_size: int = Field(ge=0)
    checksum: str
    stored_at: datetime
    tables: list[UploadedTableSummary]


def _kind(value: object) -> str:
    if value is None or value == "":
        return "empty"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)):
        return "number"
    if hasattr(value, "isoformat") and not isinstance(value, str):
        return "date"
    if isinstance(value, str):
        stripped = value.strip()
        if not stripped:
            return "empty"
        try:
            float(stripped.replace(",", ""))
            return "number"
        except ValueError:
            return "string"
    return "string"


def _summarize_columns(names: list[str], rows: list[list[object]]) -> list[UploadedColumnSummary]:
    output: list[UploadedColumnSummary] = []
    for index, name in enumerate(names):
        kinds: Counter[str] = Counter()
        for row in rows:
            kinds[_kind(row[index] if index < len(row) else None)] += 1
        non_empty = {kind for kind in kinds if kind != "empty"}
        inferred = (
            "empty" if not non_empty else next(iter(non_empty)) if len(non_empty) == 1 else "mixed"
        )
        output.append(
            UploadedColumnSummary(
                name=name,
                inferred_type=inferred,  # type: ignore[arg-type]
                non_null_count=sum(count for kind, count in kinds.items() if kind != "empty"),
                null_count=kinds["empty"],
            )
        )
    return output


def _inspect_csv(path: Path) -> list[UploadedTableSummary]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        records = list(csv.reader(stream))
    if not records:
        raise UploadValidationError("CSV file is empty.")
    names = [value.strip() or f"column_{index + 1}" for index, value in enumerate(records[0])]
    rows: list[list[object]] = [list(row) for row in records[1:]]
    return [
        UploadedTableSummary(
            name="Data",
            row_count=len(rows),
            column_count=len(names),
            columns=_summarize_columns(names, rows),
        )
    ]


def _inspect_json(path: Path) -> list[UploadedTableSummary]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(payload, dict):
        for candidate in ("data", "records", "observations"):
            if isinstance(payload.get(candidate), list):
                payload = payload[candidate]
                break
        else:
            payload = [payload]
    if not isinstance(payload, list) or any(not isinstance(item, dict) for item in payload):
        raise UploadValidationError("JSON must contain an object or an array of objects.")
    names = list(dict.fromkeys(str(key) for item in payload for key in item))
    rows = [[item.get(name) for name in names] for item in payload]
    return [
        UploadedTableSummary(
            name="Data",
            row_count=len(rows),
            column_count=len(names),
            columns=_summarize_columns(names, rows),
        )
    ]


def _inspect_xlsx(path: Path) -> list[UploadedTableSummary]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    tables: list[UploadedTableSummary] = []
    try:
        for sheet in workbook.worksheets:
            values = [list(row) for row in sheet.iter_rows(values_only=True)]
            if not values:
                continue
            names = [
                str(value).strip() if value not in (None, "") else f"column_{index + 1}"
                for index, value in enumerate(values[0])
            ]
            rows = values[1:]
            tables.append(
                UploadedTableSummary(
                    name=sheet.title,
                    row_count=len(rows),
                    column_count=len(names),
                    columns=_summarize_columns(names, rows),
                )
            )
    finally:
        workbook.close()
    if not tables:
        raise UploadValidationError("XLSX workbook has no non-empty worksheets.")
    return tables


class LocalUploadService:
    allowed_extensions = {".csv", ".json", ".xlsx"}

    def __init__(self, workflow_root: Path, max_upload_mb: int) -> None:
        self.workflow_root = workflow_root
        self.max_bytes = max_upload_mb * 1024 * 1024

    def store(
        self, session_id: str, filename: str, stream: BinaryIO, media_type: str
    ) -> UploadedFileRecord:
        safe_name = Path(filename).name
        extension = Path(safe_name).suffix.lower()
        if extension not in self.allowed_extensions:
            raise UploadValidationError("Only CSV, JSON and XLSX files are supported.")
        upload_id = f"upload-{uuid4().hex}"
        directory = self.workflow_root / "sessions" / session_id / "uploads" / upload_id
        directory.mkdir(parents=True, exist_ok=False)
        target = directory / f"source{extension}"
        digest = hashlib.sha256()
        size = 0
        try:
            with target.open("wb") as output:
                while chunk := stream.read(1024 * 1024):
                    size += len(chunk)
                    if size > self.max_bytes:
                        raise UploadValidationError("Upload exceeds the configured size limit.")
                    digest.update(chunk)
                    output.write(chunk)
            inspectors = {".csv": _inspect_csv, ".json": _inspect_json, ".xlsx": _inspect_xlsx}
            tables = inspectors[extension](target)
        except (
            BadZipFile,
            InvalidFileException,
            UnicodeDecodeError,
            json.JSONDecodeError,
        ) as error:
            shutil.rmtree(directory, ignore_errors=True)
            raise UploadValidationError(
                f"{extension.removeprefix('.').upper()} content could not be parsed."
            ) from error
        except Exception:
            shutil.rmtree(directory, ignore_errors=True)
            raise
        return UploadedFileRecord(
            upload_id=upload_id,
            session_id=session_id,
            filename=safe_name,
            media_type=media_type or "application/octet-stream",
            byte_size=size,
            checksum=f"sha256:{digest.hexdigest()}",
            stored_at=datetime.now(timezone.utc),
            tables=tables,
        )
