from __future__ import annotations

import hashlib
import json
from collections import Counter
from pathlib import Path
from typing import Any, Literal

import pandas as pd  # type: ignore[import-untyped]
from openpyxl import load_workbook  # type: ignore[import-untyped]
from pydantic import BaseModel, ConfigDict, Field


def _stable_hash(value: Any) -> str:
    encoded = json.dumps(
        value,
        ensure_ascii=True,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    ).encode("utf-8")
    return f"sha256:{hashlib.sha256(encoded).hexdigest()}"


def _safe_value(value: Any) -> Any:
    if value is None or bool(pd.isna(value)):
        return None
    if hasattr(value, "item"):
        return value.item()
    return value


class SpecificationChange(BaseModel):
    model_config = ConfigDict(frozen=True, extra="forbid")

    field: str
    before: Any = None
    after: Any = None


class ArtifactValueComparison(BaseModel):
    model_config = ConfigDict(frozen=True, extra="forbid")

    status: Literal["unchanged", "changed", "unavailable"]
    source: Literal["csv", "workbook", "specification", "unavailable"]
    baseline_value_hash: str | None = None
    current_value_hash: str | None = None
    baseline_row_count: int | None = None
    current_row_count: int | None = None
    baseline_column_count: int | None = None
    current_column_count: int | None = None
    added_row_count: int | None = None
    removed_row_count: int | None = None
    changed_sheets: list[str] = Field(default_factory=list)


class ArtifactComparisonReport(BaseModel):
    model_config = ConfigDict(frozen=True, extra="forbid")

    schema_version: Literal["mra-artifact-comparison-v1"] = "mra-artifact-comparison-v1"
    baseline_artifact_id: str
    current_artifact_id: str
    artifact_format: str
    status: Literal["unchanged", "presentation_changed", "data_changed", "incompatible"]
    baseline_hash: str
    current_hash: str
    hash_changed: bool
    specification_changes: list[SpecificationChange]
    value_comparison: ArtifactValueComparison


def _specification_changes(
    baseline: dict[str, Any], current: dict[str, Any]
) -> list[SpecificationChange]:
    return [
        SpecificationChange(field=field, before=baseline.get(field), after=current.get(field))
        for field in sorted(set(baseline) | set(current))
        if baseline.get(field) != current.get(field)
    ]


def _csv_values(path: Path) -> tuple[str, int, int, Counter[str]]:
    frame = pd.read_csv(path)
    rows = [
        [_safe_value(value) for value in row] for row in frame.itertuples(index=False, name=None)
    ]
    row_hashes = Counter(_stable_hash(row) for row in rows)
    return (
        _stable_hash({"columns": [str(value) for value in frame.columns], "rows": rows}),
        len(frame),
        len(frame.columns),
        row_hashes,
    )


def _workbook_values(path: Path) -> tuple[str, int, int, dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheet_hashes: dict[str, str] = {}
    row_count = 0
    column_count = 0
    for sheet in workbook.worksheets:
        rows = [[_safe_value(value) for value in row] for row in sheet.iter_rows(values_only=True)]
        sheet_hashes[sheet.title] = _stable_hash(rows)
        row_count += max(len(rows) - 1, 0)
        column_count = max(column_count, max((len(row) for row in rows), default=0))
    workbook.close()
    return _stable_hash(sheet_hashes), row_count, column_count, sheet_hashes


def _compare_values(
    artifact_format: str,
    baseline_path: Path,
    current_path: Path,
    baseline_specification: dict[str, Any],
    current_specification: dict[str, Any],
) -> ArtifactValueComparison:
    if artifact_format == "csv":
        before_hash, before_rows, before_columns, before_row_hashes = _csv_values(baseline_path)
        after_hash, after_rows, after_columns, after_row_hashes = _csv_values(current_path)
        return ArtifactValueComparison(
            status="unchanged" if before_hash == after_hash else "changed",
            source="csv",
            baseline_value_hash=before_hash,
            current_value_hash=after_hash,
            baseline_row_count=before_rows,
            current_row_count=after_rows,
            baseline_column_count=before_columns,
            current_column_count=after_columns,
            added_row_count=sum((after_row_hashes - before_row_hashes).values()),
            removed_row_count=sum((before_row_hashes - after_row_hashes).values()),
        )
    if artifact_format == "xlsx":
        before_hash, before_rows, before_columns, before_sheets = _workbook_values(baseline_path)
        after_hash, after_rows, after_columns, after_sheets = _workbook_values(current_path)
        changed_sheets = sorted(
            name
            for name in set(before_sheets) | set(after_sheets)
            if before_sheets.get(name) != after_sheets.get(name)
        )
        return ArtifactValueComparison(
            status="unchanged" if before_hash == after_hash else "changed",
            source="workbook",
            baseline_value_hash=before_hash,
            current_value_hash=after_hash,
            baseline_row_count=before_rows,
            current_row_count=after_rows,
            baseline_column_count=before_columns,
            current_column_count=after_columns,
            changed_sheets=changed_sheets,
        )
    specification_before_hash = baseline_specification.get("data_hash")
    specification_after_hash = current_specification.get("data_hash")
    if specification_before_hash and specification_after_hash:
        return ArtifactValueComparison(
            status=(
                "unchanged" if specification_before_hash == specification_after_hash else "changed"
            ),
            source="specification",
            baseline_value_hash=str(specification_before_hash),
            current_value_hash=str(specification_after_hash),
        )
    return ArtifactValueComparison(status="unavailable", source="unavailable")


def compare_artifacts(
    *,
    baseline_artifact_id: str,
    current_artifact_id: str,
    baseline_format: str,
    current_format: str,
    baseline_hash: str,
    current_hash: str,
    baseline_path: Path,
    current_path: Path,
    baseline_specification: dict[str, Any] | None = None,
    current_specification: dict[str, Any] | None = None,
) -> ArtifactComparisonReport:
    baseline_specification = baseline_specification or {}
    current_specification = current_specification or {}
    spec_changes = _specification_changes(baseline_specification, current_specification)
    if baseline_format != current_format:
        values = ArtifactValueComparison(status="unavailable", source="unavailable")
        status: Literal["unchanged", "presentation_changed", "data_changed", "incompatible"] = (
            "incompatible"
        )
    else:
        values = _compare_values(
            current_format,
            baseline_path,
            current_path,
            baseline_specification,
            current_specification,
        )
        if values.status == "changed":
            status = "data_changed"
        elif baseline_hash != current_hash or spec_changes:
            status = "presentation_changed"
        else:
            status = "unchanged"
    return ArtifactComparisonReport(
        baseline_artifact_id=baseline_artifact_id,
        current_artifact_id=current_artifact_id,
        artifact_format=current_format,
        status=status,
        baseline_hash=baseline_hash,
        current_hash=current_hash,
        hash_changed=baseline_hash != current_hash,
        specification_changes=spec_changes,
        value_comparison=values,
    )
