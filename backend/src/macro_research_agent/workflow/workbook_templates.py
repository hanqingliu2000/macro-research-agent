from __future__ import annotations

from pathlib import Path
from typing import Literal

from openpyxl import load_workbook  # type: ignore[import-untyped]
from pydantic import BaseModel, ConfigDict, Field

WORKBOOK_TEMPLATE_VERSION = "research-workbook-v1"


class WorkbookTemplateDefinition(BaseModel):
    model_config = ConfigDict(frozen=True, extra="forbid")

    version: Literal["research-workbook-v1"] = "research-workbook-v1"
    status: Literal["approved"] = "approved"
    guideline_version: Literal["publication-v1"] = "publication-v1"
    required_sheets: tuple[str, ...] = (
        "Analytical Data",
        "Model Results",
        "Provenance",
        "Diagnostics",
        "Chart Data",
    )
    required_headers: dict[str, tuple[str, ...]] = Field(
        default_factory=lambda: {
            "Analytical Data": ("entity_code", "indicator_code", "period", "value"),
            "Model Results": (
                "job_id",
                "model",
                "observations",
                "dropped",
                "result_hash",
                "result_json",
            ),
            "Provenance": ("field", "value"),
            "Diagnostics": ("level", "message"),
        }
    )
    frozen_header_sheets: tuple[str, ...] = (
        "Analytical Data",
        "Model Results",
        "Provenance",
        "Diagnostics",
        "Chart Data",
    )
    formula_free_sheets: tuple[str, ...] = ("Analytical Data", "Chart Data")


class WorkbookLintFinding(BaseModel):
    model_config = ConfigDict(frozen=True, extra="forbid")

    rule_id: str
    severity: Literal["warning", "fail"] = "fail"
    message: str
    location: str | None = None
    blocking: bool = True


class WorkbookLintReport(BaseModel):
    model_config = ConfigDict(frozen=True, extra="forbid")

    template_version: str
    guideline_version: str
    findings: list[WorkbookLintFinding]

    @property
    def valid(self) -> bool:
        return not any(item.blocking for item in self.findings)


WORKBOOK_TEMPLATE = WorkbookTemplateDefinition()


def get_workbook_template(version: str) -> WorkbookTemplateDefinition:
    if version != WORKBOOK_TEMPLATE_VERSION:
        raise ValueError(f"Workbook template is not registered: {version}.")
    return WORKBOOK_TEMPLATE


def lint_workbook(
    path: Path,
    *,
    template_version: str = WORKBOOK_TEMPLATE_VERSION,
    required_optional_sheets: set[str] | None = None,
    require_chart: bool = False,
) -> WorkbookLintReport:
    template = get_workbook_template(template_version)
    findings: list[WorkbookLintFinding] = []
    try:
        workbook = load_workbook(path, data_only=False, read_only=False)
    except Exception as error:
        return WorkbookLintReport(
            template_version=template.version,
            guideline_version=template.guideline_version,
            findings=[
                WorkbookLintFinding(
                    rule_id="WORKBOOK-OPEN-001",
                    message=f"Workbook could not be opened: {type(error).__name__}.",
                    location=path.name,
                )
            ],
        )

    required_sheets = set(template.required_sheets) | (required_optional_sheets or set())
    missing_sheets = sorted(required_sheets - set(workbook.sheetnames))
    if missing_sheets:
        findings.append(
            WorkbookLintFinding(
                rule_id="WORKBOOK-SHEETS-001",
                message=f"Required sheets are missing: {', '.join(missing_sheets)}.",
            )
        )

    available_base_order = [
        name for name in workbook.sheetnames if name in template.required_sheets
    ]
    expected_base_order = [name for name in template.required_sheets if name in workbook.sheetnames]
    if available_base_order != expected_base_order:
        findings.append(
            WorkbookLintFinding(
                rule_id="WORKBOOK-SHEET-ORDER-001",
                severity="warning",
                message="Core workbook sheets are not in the approved order.",
                blocking=False,
            )
        )

    for sheet_name, required_headers in template.required_headers.items():
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        headers = {str(cell.value) for cell in sheet[1] if cell.value is not None}
        missing_headers = sorted(set(required_headers) - headers)
        if missing_headers:
            findings.append(
                WorkbookLintFinding(
                    rule_id="WORKBOOK-COLUMNS-001",
                    message=f"Required columns are missing: {', '.join(missing_headers)}.",
                    location=sheet_name,
                )
            )

    for sheet_name in template.frozen_header_sheets:
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        if sheet.freeze_panes != "A2":
            findings.append(
                WorkbookLintFinding(
                    rule_id="WORKBOOK-FREEZE-001",
                    message="The header row must be frozen at A2.",
                    location=sheet_name,
                )
            )
        if sheet.max_row > 1 and not sheet.auto_filter.ref:
            findings.append(
                WorkbookLintFinding(
                    rule_id="WORKBOOK-FILTER-001",
                    message="A populated table must expose header filters.",
                    location=sheet_name,
                )
            )

    for sheet_name in template.formula_free_sheets:
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        formula_cell = next(
            (cell.coordinate for row in sheet.iter_rows() for cell in row if cell.data_type == "f"),
            None,
        )
        if formula_cell:
            findings.append(
                WorkbookLintFinding(
                    rule_id="WORKBOOK-FORMULA-001",
                    message="Published analytical values must not contain formulas.",
                    location=f"{sheet_name}!{formula_cell}",
                )
            )

    if workbook.properties.description != template.guideline_version:
        findings.append(
            WorkbookLintFinding(
                rule_id="WORKBOOK-GUIDELINE-001",
                message="Workbook metadata is not bound to the approved publication guideline.",
                location="workbook.properties.description",
            )
        )
    if getattr(workbook, "_external_links", []):
        findings.append(
            WorkbookLintFinding(
                rule_id="WORKBOOK-EXTERNAL-LINK-001",
                message="External workbook links are not allowed in a governed output.",
            )
        )
    if require_chart:
        chart_count = sum(len(sheet._charts) for sheet in workbook.worksheets)
        if chart_count < 1:
            findings.append(
                WorkbookLintFinding(
                    rule_id="WORKBOOK-CHART-001",
                    message="The selected chart deliverable is missing from the workbook.",
                    location="Chart Data",
                )
            )
    workbook.close()
    return WorkbookLintReport(
        template_version=template.version,
        guideline_version=template.guideline_version,
        findings=findings,
    )
