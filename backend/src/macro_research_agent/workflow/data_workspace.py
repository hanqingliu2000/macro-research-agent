from __future__ import annotations

import hashlib
import io
import json
import sqlite3
import threading
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Literal
from uuid import uuid4

import pandas as pd  # type: ignore[import-untyped]
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from openpyxl.worksheet.table import Table, TableStyleInfo  # type: ignore[import-untyped]
from pydantic import BaseModel, Field

from macro_research_agent.workflow.uploads import UploadedFileRecord

DataOperationKind = Literal[
    "filter_missing",
    "keep_columns",
    "drop_columns",
    "rename_columns",
    "map_codes",
    "cast_types",
    "missing_value_rule",
    "deduplicate",
    "sort",
]
DATA_OPERATION_KINDS = {
    "filter_missing",
    "keep_columns",
    "drop_columns",
    "rename_columns",
    "map_codes",
    "cast_types",
    "missing_value_rule",
    "deduplicate",
    "sort",
}

CANONICAL_KEYS = ["entity_code", "indicator_code", "frequency", "period"]
CANONICAL_COLUMNS = ["series_id", *CANONICAL_KEYS, "value"]
ROLE_ALIASES = {
    "entity_code": {"entity_code", "country_code", "country", "economy", "iso3", "ref_area"},
    "indicator_code": {"indicator_code", "indicator", "series", "measure", "concept"},
    "frequency": {"frequency", "freq"},
    "period": {"period", "date", "year", "time", "time_period"},
    "value": {"value", "obs_value", "observation", "amount"},
}
ENTITY_CODE_MAP = {
    "UNITED STATES": "USA",
    "US": "USA",
    "U.S.": "USA",
    "JAPAN": "JPN",
    "GERMANY": "DEU",
    "FRANCE": "FRA",
    "ITALY": "ITA",
    "CANADA": "CAN",
    "UNITED KINGDOM": "GBR",
    "UK": "GBR",
}


class DataWorkspaceError(ValueError):
    pass


class DatasetSnapshotRecord(BaseModel):
    session_id: str
    snapshot_id: str
    snapshot_hash: str
    schema_id: str
    label: str
    source_kind: Literal["connector", "prepared", "external_round_trip", "analytical"]
    source_ref: str
    parent_snapshot_id: str | None = None
    path: str
    row_count: int = Field(ge=0)
    columns: list[dict[str, Any]]
    created_at: datetime


class DataOperationRecord(BaseModel):
    operation_id: str
    session_id: str
    kind: DataOperationKind
    label: str
    input_snapshot_id: str
    output_snapshot_id: str
    parameters: dict[str, Any] = Field(default_factory=dict)
    created_at: datetime


class RoundTripPackage(BaseModel):
    package_id: str
    session_id: str
    source_snapshot_id: str
    schema_id: str
    path: str
    filename: str
    workbook_checksum: str
    manifest_checksum: str
    created_at: datetime


class RoundTripImportResult(BaseModel):
    import_id: str
    filename: str
    workbook_checksum: str
    manifest_verified: bool
    formulas_materialized: bool
    source_columns: list[str]
    mappings: dict[str, str]
    errors: list[str]
    warnings: list[str]
    source_snapshot_id: str | None = None
    snapshot: DatasetSnapshotRecord | None = None


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _identifier(prefix: str) -> str:
    return f"{prefix}-{uuid4().hex}"


def _sha256_bytes(value: bytes) -> str:
    return f"sha256:{hashlib.sha256(value).hexdigest()}"


def _canonical_hash(value: Any) -> str:
    payload = json.dumps(value, ensure_ascii=False, separators=(",", ":"), sort_keys=True)
    return _sha256_bytes(payload.encode())


def _safe_rows(frame: pd.DataFrame, limit: int = 100) -> list[dict[str, Any]]:
    safe = frame.head(limit).copy()
    safe = safe.astype(object).where(pd.notna(safe), None)
    return [{str(key): value for key, value in row.items()} for row in safe.to_dict("records")]


def _column_wire(name: str) -> dict[str, Any]:
    if name == "value":
        data_type = "number"
    elif name == "period":
        data_type = "period"
    else:
        data_type = "string"
    role = name if name in CANONICAL_COLUMNS else "attribute"
    return {
        "id": name,
        "label": name,
        "dataType": data_type,
        "role": role,
        "required": name in CANONICAL_COLUMNS,
        "key": name in CANONICAL_KEYS,
    }


def _normalize_name(value: object) -> str:
    return str(value).strip().lower().replace(" ", "_").replace("-", "_")


def _infer_frequency(period: str) -> str:
    upper = period.upper()
    if "Q" in upper:
        return "quarterly"
    if len(period) >= 7 and period[4:5] in {"-", "/"}:
        return "monthly"
    return "annual"


def _entity_code(value: object) -> str:
    normalized = str(value).strip().upper()
    return ENTITY_CODE_MAP.get(normalized, normalized)


def _indicator_code(value: object) -> str:
    return _normalize_name(value).upper()


def _role_source(frame: pd.DataFrame, role: str, mapping: dict[str, str] | None) -> str | None:
    source_names = {str(column): _normalize_name(column) for column in frame.columns}
    explicit = next((source for source, target in (mapping or {}).items() if target == role), None)
    if explicit in source_names:
        return explicit
    return next(
        (source for source, normalized in source_names.items() if normalized in ROLE_ALIASES[role]),
        None,
    )


class DataWorkspaceManager:
    def __init__(self, database_path: Path, workflow_root: Path) -> None:
        self.database_path = database_path
        self.workflow_root = workflow_root
        self._lock = threading.RLock()
        self._initialize()

    def _connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.database_path, timeout=30)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")
        return connection

    def _initialize(self) -> None:
        with self._connect() as connection:
            connection.executescript(
                """
                CREATE TABLE IF NOT EXISTS workflow_dataset_snapshots (
                    session_id TEXT NOT NULL,
                    snapshot_id TEXT NOT NULL,
                    record_json TEXT NOT NULL,
                    PRIMARY KEY (session_id, snapshot_id),
                    FOREIGN KEY (session_id) REFERENCES workflow_sessions(session_id)
                );
                CREATE TABLE IF NOT EXISTS workflow_data_operations (
                    session_id TEXT NOT NULL,
                    operation_id TEXT NOT NULL,
                    record_json TEXT NOT NULL,
                    PRIMARY KEY (session_id, operation_id),
                    FOREIGN KEY (session_id) REFERENCES workflow_sessions(session_id)
                );
                CREATE TABLE IF NOT EXISTS workflow_roundtrip_packages (
                    session_id TEXT NOT NULL,
                    package_id TEXT NOT NULL,
                    record_json TEXT NOT NULL,
                    PRIMARY KEY (session_id, package_id),
                    FOREIGN KEY (session_id) REFERENCES workflow_sessions(session_id)
                );
                CREATE TABLE IF NOT EXISTS workflow_roundtrip_imports (
                    session_id TEXT NOT NULL,
                    import_id TEXT NOT NULL,
                    path TEXT NOT NULL,
                    PRIMARY KEY (session_id, import_id),
                    FOREIGN KEY (session_id) REFERENCES workflow_sessions(session_id)
                );
                """
            )

    def _upload_path(self, record: UploadedFileRecord) -> Path:
        extension = Path(record.filename).suffix.lower()
        return (
            self.workflow_root
            / "sessions"
            / record.session_id
            / "uploads"
            / record.upload_id
            / f"source{extension}"
        )

    def _read_upload(self, record: UploadedFileRecord) -> pd.DataFrame:
        path = self._upload_path(record)
        extension = path.suffix.lower()
        if extension == ".csv":
            return pd.read_csv(path)
        if extension == ".json":
            payload = json.loads(path.read_text(encoding="utf-8"))
            if isinstance(payload, dict):
                payload = next(
                    (
                        payload[key]
                        for key in ("data", "records", "observations")
                        if isinstance(payload.get(key), list)
                    ),
                    [payload],
                )
            return pd.DataFrame(payload)
        if extension == ".xlsx":
            return pd.read_excel(path, sheet_name=0, engine="openpyxl")
        raise DataWorkspaceError("Unsupported uploaded file type.")

    def read_upload_frame(self, record: UploadedFileRecord) -> pd.DataFrame:
        """Read actual uploaded values for controlled local execution or model-data grants."""
        path = self._upload_path(record)
        if _sha256_bytes(path.read_bytes()) != record.checksum:
            raise DataWorkspaceError("The uploaded source checksum no longer matches its record.")
        return self._read_upload(record)

    @staticmethod
    def _normalize(
        frame: pd.DataFrame,
        *,
        mapping: dict[str, str] | None = None,
        default_entity: str | None = None,
    ) -> pd.DataFrame:
        if frame.empty:
            raise DataWorkspaceError("The uploaded table contains no observations.")
        source_names = {str(column): _normalize_name(column) for column in frame.columns}
        role_sources: dict[str, str] = {}
        for source, role in (mapping or {}).items():
            if source in source_names and role in ROLE_ALIASES:
                role_sources[role] = source
        for role, aliases in ROLE_ALIASES.items():
            if role not in role_sources:
                match = next(
                    (
                        source
                        for source, normalized in source_names.items()
                        if normalized in aliases
                    ),
                    None,
                )
                if match:
                    role_sources[role] = match

        period_source = role_sources.get("period")
        if period_source is None:
            raise DataWorkspaceError("Map one uploaded column to the required period role.")
        entity_source = role_sources.get("entity_code")
        if entity_source is None and default_entity is None:
            raise DataWorkspaceError("Map one uploaded column to the required entity_code role.")

        value_source = role_sources.get("value")
        indicator_source = role_sources.get("indicator_code")
        frequency_source = role_sources.get("frequency")
        used = {value for value in role_sources.values()}
        if value_source:
            normalized = pd.DataFrame(
                {
                    "entity_code": frame[entity_source] if entity_source else default_entity,
                    "indicator_code": frame[indicator_source]
                    if indicator_source
                    else _normalize_name(value_source).upper(),
                    "period": frame[period_source],
                    "value": frame[value_source],
                }
            )
            if frequency_source:
                normalized["frequency"] = frame[frequency_source]
            attribute_sources = [
                column
                for column in frame.columns
                if str(column) not in used and _normalize_name(column) != "series_id"
            ]
            for source in attribute_sources:
                normalized[_normalize_name(source)] = frame[source]
        else:
            candidates = [
                str(column)
                for column in frame.columns
                if str(column) not in used and pd.api.types.is_numeric_dtype(frame[column])
            ]
            if not candidates:
                raise DataWorkspaceError(
                    "Map one numeric uploaded column to the required value role."
                )
            id_vars = [
                source for source in (entity_source, period_source, frequency_source) if source
            ]
            melted = frame.melt(
                id_vars=id_vars,
                value_vars=candidates,
                var_name="indicator_code",
                value_name="value",
            )
            normalized = pd.DataFrame(
                {
                    "entity_code": melted[entity_source] if entity_source else default_entity,
                    "indicator_code": melted["indicator_code"],
                    "period": melted[period_source],
                    "value": melted["value"],
                }
            )
            if frequency_source:
                normalized["frequency"] = melted[frequency_source]

        normalized["entity_code"] = normalized["entity_code"].map(_entity_code)
        normalized["indicator_code"] = normalized["indicator_code"].map(_indicator_code)
        normalized["period"] = normalized["period"].astype(str).str.replace(r"\.0$", "", regex=True)
        normalized["value"] = pd.to_numeric(normalized["value"], errors="coerce")
        if "frequency" not in normalized:
            normalized["frequency"] = normalized["period"].map(_infer_frequency)
        else:
            frequency_map = {
                "a": "annual",
                "annual": "annual",
                "yearly": "annual",
                "q": "quarterly",
                "quarterly": "quarterly",
                "m": "monthly",
                "monthly": "monthly",
            }
            normalized["frequency"] = (
                normalized["frequency"].astype(str).str.strip().str.lower().map(frequency_map)
            )
            if normalized["frequency"].isna().any():
                raise DataWorkspaceError(
                    "Frequency values must resolve to annual, quarterly or monthly."
                )
        normalized.insert(
            0,
            "series_id",
            normalized["indicator_code"].astype(str) + ":" + normalized["entity_code"].astype(str),
        )
        ordered = [
            *CANONICAL_COLUMNS,
            *[column for column in normalized if column not in CANONICAL_COLUMNS],
        ]
        return normalized[ordered].reset_index(drop=True)

    def _snapshot_path(self, session_id: str, snapshot_id: str) -> Path:
        return (
            self.workflow_root
            / "sessions"
            / session_id
            / "data"
            / "snapshots"
            / f"{snapshot_id}.parquet"
        )

    def _record_snapshot(
        self,
        session_id: str,
        frame: pd.DataFrame,
        *,
        label: str,
        source_kind: Literal["connector", "prepared", "external_round_trip", "analytical"],
        source_ref: str,
        parent_snapshot_id: str | None = None,
    ) -> DatasetSnapshotRecord:
        snapshot_id = _identifier("dataset")
        path = self._snapshot_path(session_id, snapshot_id)
        path.parent.mkdir(parents=True, exist_ok=True)
        buffer = io.BytesIO()
        frame.to_parquet(buffer, index=False)
        payload = buffer.getvalue()
        path.write_bytes(payload)
        columns = [_column_wire(str(column)) for column in frame.columns]
        schema_id = _canonical_hash(columns)
        record = DatasetSnapshotRecord(
            session_id=session_id,
            snapshot_id=snapshot_id,
            snapshot_hash=_sha256_bytes(payload),
            schema_id=schema_id,
            label=label,
            source_kind=source_kind,
            source_ref=source_ref,
            parent_snapshot_id=parent_snapshot_id,
            path=str(path),
            row_count=len(frame),
            columns=columns,
            created_at=_now(),
        )
        with self._lock, self._connect() as connection:
            connection.execute(
                """
                INSERT INTO workflow_dataset_snapshots(session_id, snapshot_id, record_json)
                VALUES (?, ?, ?)
                """,
                (session_id, snapshot_id, record.model_dump_json()),
            )
        return record

    def record_analytical_snapshot(
        self,
        session_id: str,
        frame: pd.DataFrame,
        *,
        label: str,
        operation_id: str,
        parent_snapshot_id: str,
    ) -> DatasetSnapshotRecord:
        return self._record_snapshot(
            session_id,
            frame,
            label=label,
            source_kind="analytical",
            source_ref=operation_id,
            parent_snapshot_id=parent_snapshot_id,
        )

    def get_snapshot(self, session_id: str, snapshot_id: str) -> DatasetSnapshotRecord:
        with self._connect() as connection:
            row = connection.execute(
                """
                SELECT record_json FROM workflow_dataset_snapshots
                WHERE session_id = ? AND snapshot_id = ?
                """,
                (session_id, snapshot_id),
            ).fetchone()
        if row is None:
            raise DataWorkspaceError("Unknown dataset snapshot.")
        return DatasetSnapshotRecord.model_validate_json(row["record_json"])

    def frame(self, record: DatasetSnapshotRecord) -> pd.DataFrame:
        return pd.read_parquet(record.path)

    def preview(
        self, record: DatasetSnapshotRecord, offset: int = 0, limit: int = 100
    ) -> list[dict[str, Any]]:
        if offset < 0 or limit < 1 or limit > 500:
            raise DataWorkspaceError("Preview offset/limit is outside the allowed range.")
        frame = self.frame(record).iloc[offset : offset + limit]
        return _safe_rows(frame, limit)

    def diagnostics(self, record: DatasetSnapshotRecord) -> dict[str, Any]:
        frame = self.frame(record)
        observation_count = len(frame)
        missing_observations = int(frame["value"].isna().sum()) if "value" in frame else 0
        duplicates = int(frame.duplicated(CANONICAL_KEYS, keep=False).sum())
        warnings: list[str] = []
        if missing_observations:
            warnings.append(
                f"{missing_observations} observation value(s) are missing. "
                "Review before pipeline execution."
            )
        if duplicates:
            warnings.append("Duplicate keys require review.")
        return {
            "missingShare": round(missing_observations / max(observation_count, 1) * 100, 2),
            "duplicateKeys": duplicates,
            "observations": observation_count,
            "warning": " ".join(warnings),
        }

    def validate_for_approval(self, record: DatasetSnapshotRecord) -> None:
        self.validate_frame(self.frame(record))

    def validate_frame(self, frame: pd.DataFrame) -> None:
        missing_columns = set(CANONICAL_COLUMNS) - set(frame.columns)
        if missing_columns:
            raise DataWorkspaceError(
                "Dataset is missing canonical fields: " + ", ".join(sorted(missing_columns))
            )
        if frame[CANONICAL_KEYS].isna().any().any():
            raise DataWorkspaceError("Dataset keys cannot contain missing values.")
        if frame.duplicated(CANONICAL_KEYS).any():
            raise DataWorkspaceError("Dataset contains duplicate entity-indicator-period keys.")
        if not pd.api.types.is_numeric_dtype(frame["value"]):
            raise DataWorkspaceError("Dataset value field must be numeric.")

    def snapshot_wire(self, record: DatasetSnapshotRecord) -> dict[str, Any]:
        return {
            "id": record.snapshot_id,
            "hash": record.snapshot_hash,
            "schemaId": record.schema_id,
            "label": record.label,
            "sourceKind": record.source_kind,
            "parentSnapshotId": record.parent_snapshot_id,
            "rows": self.preview(record),
            "createdAt": record.created_at.isoformat(),
        }

    def initialize_connector(self, session_id: str, result: Any) -> dict[str, Any]:
        rows = []
        for observation in result.dataset.observations:
            row = observation.model_dump(exclude={"attributes"})
            row["frequency"] = {"M": "monthly", "Q": "quarterly", "A": "annual"}.get(
                str(row["frequency"]), str(row["frequency"])
            )
            for namespace, attributes in observation.attributes.items():
                if not isinstance(attributes, dict):
                    continue
                row.update(
                    {
                        f"{namespace}_{key.lower()}": value
                        for key, value in attributes.items()
                        if not isinstance(value, (dict, list))
                    }
                )
            rows.append(row)
        frame = pd.DataFrame(rows)
        self.validate_frame(frame)
        provenance = result.provenance.model_dump(mode="json")
        record = self._record_snapshot(
            session_id,
            frame,
            label=f"{result.provenance.provider} {result.provenance.dataset}",
            source_kind="connector",
            source_ref=result.provenance.raw_sha256,
        )
        return {
            "schemaId": record.schema_id,
            "columns": record.columns,
            "sourceRuns": [
                {
                    "id": _identifier("source-run"),
                    "provider": result.provenance.provider,
                    "dataset": result.provenance.dataset,
                    "connectorId": result.provenance.connector_id,
                    "query": result.provenance.source_identifier,
                    "status": "completed",
                    "retrievedAt": result.provenance.retrieved_at.isoformat(),
                }
            ],
            "snapshots": [self.snapshot_wire(record)],
            "activeSnapshotId": record.snapshot_id,
            "operations": [],
            "quality": self.diagnostics(record),
            "entityMappings": {str(code): str(code) for code in frame["entity_code"].unique()},
            "indicatorMappings": {
                str(code): str(code) for code in frame["indicator_code"].unique()
            },
            "provenance": [
                {
                    "sourceKind": result.provenance.connector_id,
                    "normalizedSnapshotId": record.snapshot_id,
                    **provenance,
                }
            ],
        }

    def merge_connector_snapshots(
        self,
        session_id: str,
        inputs: list[tuple[DatasetSnapshotRecord, dict[str, str]]],
    ) -> DatasetSnapshotRecord:
        """Create a reviewable canonical union while retaining row-level source lineage."""
        if len(inputs) < 2:
            raise DataWorkspaceError("A cross-source merge requires at least two snapshots.")
        frames: list[pd.DataFrame] = []
        for record, source in inputs:
            frame = self.frame(record).copy()
            frame["source_connector_id"] = source["connectorId"]
            frame["source_dataset"] = source["dataset"]
            frame["source_snapshot_id"] = record.snapshot_id
            frames.append(frame)
        merged = pd.concat(frames, ignore_index=True, sort=False)
        duplicate_rows = merged[merged.duplicated(CANONICAL_KEYS, keep=False)]
        if not duplicate_rows.empty:
            examples = duplicate_rows[CANONICAL_KEYS].drop_duplicates().head(3)
            rendered = "; ".join(
                " / ".join(str(row[key]) for key in CANONICAL_KEYS)
                for _, row in examples.iterrows()
            )
            raise DataWorkspaceError(
                "Cross-source observations overlap on canonical identity keys. "
                f"Resolve indicator identities before merge: {rendered}."
            )
        merged = merged.sort_values(CANONICAL_KEYS, kind="stable").reset_index(drop=True)
        self.validate_frame(merged)
        return self._record_snapshot(
            session_id,
            merged,
            label="Reviewed cross-source dataset",
            source_kind="prepared",
            source_ref="canonical_union:v1",
        )

    def initialize(
        self,
        session: dict[str, Any],
        upload: UploadedFileRecord,
        mapping: dict[str, str] | None = None,
    ) -> dict[str, Any]:
        plan = session.get("currentPlan") or {}
        resolutions = plan.get("geographyResolutions") or []
        members = resolutions[0].get("memberCodes", []) if len(resolutions) == 1 else []
        default_entity = members[0] if len(members) == 1 else None
        raw_frame = self._read_upload(upload)
        entity_source = _role_source(raw_frame, "entity_code", mapping)
        indicator_source = _role_source(raw_frame, "indicator_code", mapping)
        entity_mappings = (
            {
                str(value): _entity_code(value)
                for value in raw_frame[entity_source].dropna().unique()
            }
            if entity_source
            else ({default_entity: default_entity} if default_entity else {})
        )
        indicator_mappings = (
            {
                str(value): _indicator_code(value)
                for value in raw_frame[indicator_source].dropna().unique()
            }
            if indicator_source
            else {}
        )
        normalized = self._normalize(raw_frame, mapping=mapping, default_entity=default_entity)
        record = self._record_snapshot(
            upload.session_id,
            normalized,
            label=f"Normalized {upload.filename}",
            source_kind="connector",
            source_ref=upload.upload_id,
        )
        return {
            "schemaId": record.schema_id,
            "columns": record.columns,
            "sourceRuns": [
                {
                    "id": _identifier("source-run"),
                    "provider": "Local file",
                    "dataset": upload.filename,
                    "connectorId": "local_file",
                    "query": f"Normalize upload {upload.upload_id}",
                    "status": "completed",
                    "retrievedAt": _now().isoformat(),
                }
            ],
            "snapshots": [self.snapshot_wire(record)],
            "activeSnapshotId": record.snapshot_id,
            "operations": [],
            "quality": self.diagnostics(record),
            "entityMappings": entity_mappings,
            "indicatorMappings": indicator_mappings,
            "provenance": [
                {
                    "sourceKind": "local_file",
                    "uploadId": upload.upload_id,
                    "filename": upload.filename,
                    "rawChecksum": upload.checksum,
                    "normalizedSnapshotId": record.snapshot_id,
                    "normalizedChecksum": record.snapshot_hash,
                    "retrievedAt": upload.stored_at.isoformat(),
                }
            ],
            "previewRef": (f"/api/v1/workflow/sessions/{upload.session_id}/data/preview"),
        }

    def source_mapping_draft(self, upload: UploadedFileRecord, error: str = "") -> dict[str, Any]:
        if not upload.tables:
            raise DataWorkspaceError("Uploaded file has no inspectable tables.")
        columns = [column.name for column in upload.tables[0].columns]
        mappings: dict[str, str] = {}
        for column in columns:
            normalized = _normalize_name(column)
            mappings[column] = next(
                (role for role, aliases in ROLE_ALIASES.items() if normalized in aliases),
                "attribute",
            )
        return {
            "uploadId": upload.upload_id,
            "filename": upload.filename,
            "sourceColumns": columns,
            "mappings": mappings,
            "state": "mapping",
            "errors": [error] if error else [],
        }

    def _operation_frame(
        self, frame: pd.DataFrame, kind: DataOperationKind, parameters: dict[str, Any]
    ) -> pd.DataFrame:
        output = frame.copy()
        if kind == "filter_missing":
            output = output[output[parameters.get("column", "value")].notna()]
        elif kind == "keep_columns":
            requested = parameters.get("columns") or CANONICAL_COLUMNS
            keep = [column for column in requested if column in output.columns]
            missing_required = set(CANONICAL_COLUMNS) - set(keep)
            if missing_required:
                raise DataWorkspaceError("Keep columns cannot remove canonical required fields.")
            output = output[keep]
        elif kind == "drop_columns":
            requested = parameters.get("columns") or [
                column for column in output if column not in CANONICAL_COLUMNS
            ]
            protected = set(requested) & set(CANONICAL_COLUMNS)
            if protected:
                raise DataWorkspaceError("Canonical required fields cannot be dropped.")
            output = output.drop(columns=[column for column in requested if column in output])
        elif kind == "rename_columns":
            renames = parameters.get("mapping") or {
                column: _normalize_name(column)
                for column in output
                if column not in CANONICAL_COLUMNS
            }
            if set(renames) & set(CANONICAL_COLUMNS):
                raise DataWorkspaceError("Canonical required fields cannot be renamed.")
            output = output.rename(columns=renames)
        elif kind == "map_codes":
            for column in parameters.get("columns") or ["entity_code", "indicator_code"]:
                output[column] = output[column].astype(str).str.strip().str.upper()
            output["series_id"] = output["indicator_code"] + ":" + output["entity_code"]
        elif kind == "cast_types":
            output["value"] = pd.to_numeric(output["value"], errors="coerce")
            output["period"] = output["period"].astype(str)
        elif kind == "missing_value_rule":
            tokens = parameters.get("tokens") or ["..", "NA", "N/A", "null"]
            output = output.replace(tokens, None)
            output["value"] = pd.to_numeric(output["value"], errors="coerce")
        elif kind == "deduplicate":
            output = output.drop_duplicates(CANONICAL_KEYS, keep=parameters.get("keep", "last"))
        elif kind == "sort":
            output = output.sort_values(parameters.get("columns") or CANONICAL_KEYS)
        return output.reset_index(drop=True)

    def apply_operation(
        self,
        session_id: str,
        active_snapshot_id: str,
        kind: DataOperationKind,
        parameters: dict[str, Any] | None = None,
    ) -> tuple[DatasetSnapshotRecord, DataOperationRecord]:
        if kind not in DATA_OPERATION_KINDS:
            raise DataWorkspaceError("Unknown Data operation.")
        active = self.get_snapshot(session_id, active_snapshot_id)
        parameters = parameters or {}
        try:
            output = self._operation_frame(self.frame(active), kind, parameters)
            snapshot = self._record_snapshot(
                session_id,
                output,
                label=f"{kind.replace('_', ' ').title()} result",
                source_kind="prepared",
                source_ref=kind,
                parent_snapshot_id=active.snapshot_id,
            )
        except (KeyError, TypeError, ValueError) as error:
            raise DataWorkspaceError(
                "Invalid parameters for the selected Data operation."
            ) from error
        operation = DataOperationRecord(
            operation_id=_identifier("data-operation"),
            session_id=session_id,
            kind=kind,
            label=kind.replace("_", " ").title(),
            input_snapshot_id=active.snapshot_id,
            output_snapshot_id=snapshot.snapshot_id,
            parameters=parameters,
            created_at=_now(),
        )
        with self._lock, self._connect() as connection:
            connection.execute(
                """
                INSERT INTO workflow_data_operations(session_id, operation_id, record_json)
                VALUES (?, ?, ?)
                """,
                (session_id, operation.operation_id, operation.model_dump_json()),
            )
        return snapshot, operation

    def export_roundtrip(self, session_id: str, snapshot_id: str) -> RoundTripPackage:
        snapshot = self.get_snapshot(session_id, snapshot_id)
        frame = self.frame(snapshot)
        package_id = _identifier("roundtrip")
        manifest = {
            "format_version": "mra-roundtrip-v1",
            "session_id": session_id,
            "source_snapshot_id": snapshot.snapshot_id,
            "source_snapshot_hash": snapshot.snapshot_hash,
            "schema_id": snapshot.schema_id,
            "package_id": package_id,
        }
        manifest_checksum = _canonical_hash(manifest)
        workbook = Workbook()
        data_sheet = workbook.active
        data_sheet.title = "Data"
        data_sheet.append([str(column) for column in frame.columns])
        for row in (
            frame.astype(object).where(pd.notna(frame), None).itertuples(index=False, name=None)
        ):
            data_sheet.append(list(row))
        if data_sheet.max_row >= 2 and data_sheet.max_column >= 1:
            table = Table(displayName="MRAData", ref=data_sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            data_sheet.add_table(table)
        schema_sheet = workbook.create_sheet("Schema")
        schema_sheet.append(["field", "data_type", "logical_role", "required", "key"])
        for column in snapshot.columns:
            schema_sheet.append(
                [
                    column["id"],
                    column["dataType"],
                    column["role"],
                    column["required"],
                    column["key"],
                ]
            )
        provenance = workbook.create_sheet("Provenance")
        provenance.append(["property", "value"])
        for key, value in manifest.items():
            provenance.append([key, value])
        instructions = workbook.create_sheet("Instructions")
        instructions.append(
            ["Edit values and rows in Data only. Do not use formulas. Preserve required fields."]
        )
        manifest_sheet = workbook.create_sheet("_MRA_MANIFEST")
        manifest_sheet.append(["property", "value"])
        for key, value in {**manifest, "manifest_checksum": manifest_checksum}.items():
            manifest_sheet.append([key, value])
        manifest_sheet.sheet_state = "hidden"
        for protected_sheet in (schema_sheet, provenance, instructions, manifest_sheet):
            protected_sheet.protection.sheet = True
        workbook.security.lockStructure = True
        path = (
            self.workflow_root
            / "sessions"
            / session_id
            / "data"
            / "roundtrip"
            / f"{package_id}.xlsx"
        )
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        payload = path.read_bytes()
        package = RoundTripPackage(
            package_id=package_id,
            session_id=session_id,
            source_snapshot_id=snapshot.snapshot_id,
            schema_id=snapshot.schema_id,
            path=str(path),
            filename=f"macro-research-cleaning-{snapshot.snapshot_id}.xlsx",
            workbook_checksum=_sha256_bytes(payload),
            manifest_checksum=manifest_checksum,
            created_at=_now(),
        )
        with self._lock, self._connect() as connection:
            connection.execute(
                """
                INSERT INTO workflow_roundtrip_packages(session_id, package_id, record_json)
                VALUES (?, ?, ?)
                """,
                (session_id, package_id, package.model_dump_json()),
            )
        return package

    def get_roundtrip_package(self, session_id: str, package_id: str) -> RoundTripPackage:
        with self._connect() as connection:
            row = connection.execute(
                """
                SELECT record_json FROM workflow_roundtrip_packages
                WHERE session_id = ? AND package_id = ?
                """,
                (session_id, package_id),
            ).fetchone()
        if row is None:
            raise DataWorkspaceError("Unknown round-trip package.")
        return RoundTripPackage.model_validate_json(row["record_json"])

    def import_roundtrip(
        self, session_id: str, filename: str, content: bytes
    ) -> RoundTripImportResult:
        import_id = _identifier("roundtrip-import")
        checksum = _sha256_bytes(content)
        import_path = (
            self.workflow_root
            / "sessions"
            / session_id
            / "data"
            / "roundtrip"
            / "imports"
            / f"{import_id}.xlsx"
        )
        import_path.parent.mkdir(parents=True, exist_ok=True)
        import_path.write_bytes(content)
        with self._lock, self._connect() as connection:
            connection.execute(
                """
                INSERT INTO workflow_roundtrip_imports(session_id, import_id, path)
                VALUES (?, ?, ?)
                """,
                (session_id, import_id, str(import_path)),
            )
        workbook = load_workbook(io.BytesIO(content), data_only=False)
        required_sheets = {"Data", "Schema", "Provenance", "Instructions", "_MRA_MANIFEST"}
        errors: list[str] = []
        warnings: list[str] = []
        if not required_sheets.issubset(workbook.sheetnames):
            missing = sorted(required_sheets - set(workbook.sheetnames))
            errors.append("Missing required workbook sheets: " + ", ".join(missing))
            return RoundTripImportResult(
                import_id=import_id,
                filename=Path(filename).name,
                workbook_checksum=checksum,
                manifest_verified=False,
                formulas_materialized=False,
                source_columns=[],
                mappings={},
                errors=errors,
                warnings=warnings,
            )
        manifest_values = {
            str(row[0].value): row[1].value
            for row in workbook["_MRA_MANIFEST"].iter_rows(min_row=2)
            if row[0].value
        }
        provided_checksum = str(manifest_values.pop("manifest_checksum", ""))
        manifest_verified = manifest_values.get(
            "session_id"
        ) == session_id and provided_checksum == _canonical_hash(manifest_values)
        if not manifest_verified:
            errors.append("The locked MRA manifest does not match this session or was modified.")
        data_sheet = workbook["Data"]
        formula_cells = [
            cell.coordinate
            for row in data_sheet.iter_rows()
            for cell in row
            if cell.data_type == "f"
        ]
        formulas_materialized = not formula_cells
        if formula_cells:
            errors.append("Data formulas must be converted to values before reimport.")
        rows = list(data_sheet.iter_rows(values_only=True))
        source_columns = [str(value) for value in rows[0]] if rows else []
        mappings: dict[str, str] = {}
        for column in source_columns:
            normalized_name = _normalize_name(column)
            mappings[column] = next(
                (role for role, aliases in ROLE_ALIASES.items() if normalized_name in aliases),
                "ignore" if normalized_name == "series_id" else "attribute",
            )
        source_snapshot_id = str(manifest_values.get("source_snapshot_id") or "") or None
        snapshot: DatasetSnapshotRecord | None = None
        required_roles = {"entity_code", "indicator_code", "frequency", "period", "value"}
        if not errors and required_roles.issubset(set(mappings.values())):
            validated = self.validate_roundtrip_import(
                session_id,
                import_id,
                mappings,
                checksum=checksum,
            )
            errors = validated.errors
            warnings = validated.warnings
            snapshot = validated.snapshot
        return RoundTripImportResult(
            import_id=import_id,
            filename=Path(filename).name,
            workbook_checksum=checksum,
            manifest_verified=manifest_verified,
            formulas_materialized=formulas_materialized,
            source_columns=source_columns,
            mappings=mappings,
            errors=errors,
            warnings=warnings,
            source_snapshot_id=source_snapshot_id,
            snapshot=snapshot,
        )

    def validate_roundtrip_import(
        self,
        session_id: str,
        import_id: str,
        mappings: dict[str, str],
        *,
        checksum: str | None = None,
    ) -> RoundTripImportResult:
        with self._connect() as connection:
            row = connection.execute(
                """
                SELECT path FROM workflow_roundtrip_imports
                WHERE session_id = ? AND import_id = ?
                """,
                (session_id, import_id),
            ).fetchone()
        if row is None:
            raise DataWorkspaceError("Unknown round-trip import draft.")
        path = Path(str(row["path"]))
        content = path.read_bytes()
        workbook = load_workbook(io.BytesIO(content), data_only=False)
        data_sheet = workbook["Data"]
        rows = list(data_sheet.iter_rows(values_only=True))
        source_columns = [str(value) for value in rows[0]] if rows else []
        required_roles = {"entity_code", "indicator_code", "frequency", "period", "value"}
        selected_roles = [role for role in mappings.values() if role in required_roles]
        missing_roles = required_roles - set(selected_roles)
        duplicate_roles = {role for role in selected_roles if selected_roles.count(role) > 1}
        errors: list[str] = []
        if missing_roles:
            errors.append("Missing required mappings: " + ", ".join(sorted(missing_roles)))
        if duplicate_roles:
            errors.append(
                "Logical roles may be mapped only once: " + ", ".join(sorted(duplicate_roles))
            )
        snapshot: DatasetSnapshotRecord | None = None
        if not errors:
            frame = pd.DataFrame(rows[1:], columns=source_columns)
            try:
                normalized = self._normalize(frame, mapping=mappings)
                self.validate_frame(normalized)
                snapshot = self._record_snapshot(
                    session_id,
                    normalized,
                    label=f"Reimported {path.name}",
                    source_kind="external_round_trip",
                    source_ref=checksum or _sha256_bytes(content),
                )
            except DataWorkspaceError as error:
                errors.append(str(error))
                snapshot = None
        return RoundTripImportResult(
            import_id=import_id,
            filename=path.name,
            workbook_checksum=checksum or _sha256_bytes(content),
            manifest_verified=True,
            formulas_materialized=True,
            source_columns=source_columns,
            mappings=mappings,
            errors=errors,
            warnings=[],
            snapshot=snapshot,
        )
